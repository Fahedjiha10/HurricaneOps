from __future__ import annotations

import csv
import json
import re
from copy import copy
from zipfile import ZIP_DEFLATED, ZipFile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from shutil import copy2, move

import fitz
import pdfplumber
from dateutil import parser as date_parser
from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from config import (
    INBOX_DIR,
    JOBS_DIR,
    JOBS_INDEX_PATH,
    ORGANIZED_PLAN_COMPARTMENTS,
    ORGANIZED_PLAN_SCHEMA_VERSION,
)
from export_quote import quote_rows_from_structured, write_extraction_outputs
from job_intake import IntakeError, _organize_source_pdfs, create_job_from_inbox
from municipality_lookup import lookup_municipality
from schedule_extractor import ScheduleExtractor
from schemas import ExtractionResult
from validators import HUMAN_REVIEW_THRESHOLD


WINDOW_HEADERS = (
    "mark",
    "quantity",
    "width",
    "height",
    "type",
    "material",
    "glass_type",
    "finish",
    "noa",
    "brand_product",
    "level",
    "remarks",
    "extraction_status",
    "source",
)
STOREFRONT_HEADERS = (
    "mark",
    "quantity",
    "width",
    "height",
    "type",
    "material",
    "glass_type",
    "finish",
    "noa",
    "brand_product",
    "level",
    "remarks",
    "extraction_status",
    "source",
)
GLAZING_HEADERS = (
    "mark",
    "description",
    "quantity",
    "width",
    "height",
    "material",
    "glass_type",
    "finish",
    "remarks",
    "noa",
    "brand_product",
    "level",
    "source",
)
DOOR_HEADERS = (
    "door_no",
    "location",
    "type",
    "width",
    "height",
    "thickness",
    "door_material",
    "door_finish",
    "frame_material",
    "frame_finish",
    "fire_rating",
    "noa",
    "panic_hardware",
    "remarks",
    "level",
    "quantity",
    "panels",
    "fixed_panels",
    "scope",
    "extraction_status",
    "raw_text",
    "source",
)
ISSUE_HEADERS = (
    "severity",
    "issue_type",
    "mark",
    "source_page",
    "description",
    "recommended_action",
)


def _extract_text(pdf_path: Path) -> str:
    with fitz.open(pdf_path) as document:
        return "\n".join(page.get_text() for page in document)


def _parse_glazing_requirements(text: str, source: str) -> dict[str, str]:
    match = re.search(
        r"\bU-?Factor\s*=\s*(\d+(?:\.\d+)?)\s*,?\s*SHGC\s*=\s*(\d+(?:\.\d+)?)\b",
        text,
        flags=re.IGNORECASE,
    )
    if not match:
        return {}
    return {
        "u_factor": match.group(1),
        "shgc": match.group(2),
        "source": source,
    }


def _extract_glazing_requirements(pdf_paths: list[Path]) -> dict[str, str]:
    requirements = [
        parsed
        for pdf_path in pdf_paths
        if (parsed := _parse_glazing_requirements(_extract_text(pdf_path), pdf_path.name))
    ]
    unique_values = {
        (requirement["u_factor"], requirement["shgc"]) for requirement in requirements
    }
    if len(unique_values) > 1:
        raise IntakeError(
            "Conflicting glazing U-factor or SHGC requirements were extracted. "
            "Review the indexed schedule sheets before issuing the quote."
        )
    return requirements[0] if requirements else {}


def _write_csv(path: Path, headers: tuple[str, ...], rows: list[dict[str, str]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _write_xlsx_report(
    path: Path, title: str, headers: tuple[str, ...], rows: list[dict[str, str]]
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.append(list(headers))
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max(max_length + 2, 12), 70)
    workbook.save(path)


def _extract_windows(text: str, source: str) -> list[dict[str, str]]:
    section = text.split("Window Schedule", 1)[-1].split("STOREFRONT SCHEDULE", 1)[0]
    pattern = re.compile(
        r"\n([A-Z]\d*)\n(\d+' - \d+(?: \d+/\d+)?\")\n(\d+' - \d+(?: \d+/\d+)?\")"
        r"\n([^\n]+)\n([^\n]+)\n([^\n]+)\n([^\n]+)"
    )
    return [
        {
            "mark": mark,
            "quantity": "",
            "width": width,
            "height": height,
            "type": window_type,
            "material": material,
            "glass_type": "",
            "finish": color,
            "noa": "",
            "brand_product": "",
            "level": "",
            "remarks": remarks,
            "extraction_status": "TEXT PARSED - legacy window schedule",
            "source": source,
        }
        for mark, width, height, window_type, material, color, remarks in pattern.findall(section)
    ]


def _extract_storefronts(text: str, source: str) -> list[dict[str, str]]:
    section = text.split("STOREFRONT SCHEDULE", 1)[-1]
    pattern = re.compile(
        r"\n(A\d+)\n(\d+' - \d+(?: \d+/\d+)?\")\n(\d+' - \d+(?: \d+/\d+)?\")"
        r"\n([^\n]+)\n([^\n]+)\n([^\n]+)"
    )
    return [
        {
            "mark": mark,
            "quantity": "",
            "width": width,
            "height": height,
            "type": storefront_type,
            "material": material,
            "glass_type": "",
            "finish": "",
            "noa": "",
            "brand_product": "",
            "level": "",
            "remarks": remarks,
            "extraction_status": "TEXT PARSED - legacy storefront schedule",
            "source": source,
        }
        for mark, width, height, storefront_type, material, remarks in pattern.findall(section)
    ]


def _clean_table_value(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def _normalize_dimension(value: str) -> str:
    normalized = _clean_table_value(value)
    return re.sub(r'(\d+)"\s+(\d+)$', r'\1/\2"', normalized)


def _join_notes(*parts: str) -> str:
    return "; ".join(part.strip() for part in parts if part and part.strip())


def _split_combined_door_frame_material(value: str) -> tuple[str, str]:
    parts = [part.strip() for part in _clean_table_value(value).split(",") if part.strip()]
    if len(parts) >= 2:
        return parts[0], parts[1]
    return _clean_table_value(value), ""


def _empty_door_row() -> dict[str, str]:
    return {header: "" for header in DOOR_HEADERS}


def _door_table_row(
    values: list[str], source: str, level: str
) -> dict[str, str] | None:
    if not values or not re.fullmatch(r"(?:\d{3}|G\d+)", values[0], flags=re.IGNORECASE):
        return None
    if len(values) >= 14 and values[2].isdigit():
        door_material, frame_finish = _split_combined_door_frame_material(values[9])
        row = _empty_door_row()
        row.update(
            {
                "door_no": values[0],
                "location": values[1],
                "quantity": values[2],
                "width": _normalize_dimension(values[3]),
                "height": _normalize_dimension(values[4]),
                "panels": values[5],
                "fixed_panels": values[6],
                "frame_material": values[7],
                "type": values[8],
                "door_material": door_material,
                "frame_finish": values[10] or frame_finish,
                "panic_hardware": values[11],
                "remarks": values[12],
                "level": level,
                "scope": "garage door schedule" if values[0].upper() == "G1" else "door schedule",
                "extraction_status": "TABLE PARSED - door schedule",
                "raw_text": " | ".join(value for value in values if value),
                "source": source,
            }
        )
        return row
    if len(values) >= 15 and re.fullmatch(r"\d{3}", values[0]):
        row = _empty_door_row()
        for header, value in zip(DOOR_HEADERS[:15], values[:15]):
            row[header] = value
        row.update(
            {
                "scope": "door schedule",
                "extraction_status": "TABLE PARSED - legacy door schedule",
                "raw_text": " | ".join(value for value in values if value),
                "source": source,
            }
        )
        return row
    return None


def _extract_doors(pdf_path: Path) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                level = ""
                for row in table:
                    values = [_clean_table_value(value) for value in row]
                    row_text = " ".join(value for value in values if value)
                    if re.fullmatch(r"LEVEL\s+\d+", row_text, flags=re.IGNORECASE):
                        level = row_text.upper()
                    parsed_row = _door_table_row(values, pdf_path.name, level)
                    if parsed_row:
                        rows.append(parsed_row)

    unique_rows: list[dict[str, str]] = []
    seen: set[tuple[str, ...]] = set()
    for row in rows:
        identity = tuple(row[header] for header in DOOR_HEADERS[:-3])
        if identity not in seen:
            seen.add(identity)
            unique_rows.append(row)

    parsed_numbers = {row["door_no"] for row in unique_rows}
    common_area_text = _extract_text(pdf_path).split("Common Areas Door Schedule1", 1)[-1]
    common_area_text = common_area_text.split("Unit Door Schedule", 1)[0]
    blocks = re.split(r"(?m)(?=^\d{3}\s*$)", common_area_text)
    for block in blocks:
        match = re.match(r"(?m)^(\d{3})\s*$", block)
        if not match or match.group(1) in parsed_numbers:
            continue
        raw_text = " | ".join(line.strip() for line in block.splitlines() if line.strip())
        fallback = _empty_door_row()
        fallback.update(
            {
                "door_no": match.group(1),
                "extraction_status": "REVIEW - raw text only",
                "raw_text": raw_text,
                "source": pdf_path.name,
            }
        )
        unique_rows.append(fallback)

    unique_rows.sort(key=lambda row: row["door_no"])
    return unique_rows


def _glazing_row_from_values(
    values: list[str], source: str, level: str
) -> dict[str, str] | None:
    if len(values) < 8 or not re.fullmatch(r"G\d+[A-Z]?", values[0], flags=re.IGNORECASE):
        return None
    return {
        "mark": values[0],
        "description": values[1],
        "quantity": values[2],
        "width": _normalize_dimension(values[3]),
        "height": _normalize_dimension(values[4]),
        "material": "",
        "glass_type": values[7],
        "finish": "",
        "remarks": values[5],
        "noa": values[6],
        "brand_product": values[7],
        "level": level,
        "source": source,
    }


def _glazing_rows_from_text(text: str, source: str) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    level = ""
    dimension = r'\d+\'-\d+(?:\s+\d+/\d+)?"'
    pattern = re.compile(
        rf"\b(G\d+[A-Z]?)\s+(.+?)\s+(\d+)\s+({dimension})\s+({dimension})(?:\s+(.*))?$",
        flags=re.IGNORECASE,
    )
    for line in text.splitlines():
        clean_line = _clean_table_value(line)
        if re.fullmatch(r"LEVEL\s+\d+", clean_line, flags=re.IGNORECASE):
            level = clean_line.upper()
            continue
        match = pattern.search(clean_line)
        if not match:
            continue
        mark, description, quantity, width, height, remainder = match.groups()
        noa_match = re.search(r"\bFL\s+\d+(?:\.\d+)?\b", remainder or "")
        if noa_match:
            remarks = (remainder or "")[: noa_match.start()].strip()
            noa = noa_match.group(0)
            brand_product = (remainder or "")[noa_match.end() :].strip()
        else:
            remarks = (remainder or "").strip()
            noa = ""
            brand_product = ""
        rows.append(
            {
                "mark": mark,
                "description": description,
                "quantity": quantity,
                "width": width,
                "height": height,
                "material": "",
                "glass_type": brand_product,
                "finish": "",
                "remarks": remarks,
                "noa": noa,
                "brand_product": brand_product,
                "level": level,
                "source": source,
            }
        )
    return rows


def _extract_glazing_rows(pdf_path: Path) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            page_rows: list[dict[str, str]] = []
            for table in page.extract_tables():
                if not table or not table[0] or _clean_table_value(table[0][0]) != "TAG":
                    continue
                level = ""
                for raw_row in table[1:]:
                    values = [_clean_table_value(value) for value in raw_row]
                    if values and re.fullmatch(r"LEVEL\s+\d+", values[0], flags=re.IGNORECASE):
                        level = values[0].upper()
                        continue
                    parsed_row = _glazing_row_from_values(values, pdf_path.name, level)
                    if parsed_row:
                        page_rows.append(parsed_row)
            page_text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
            table_marks = {row["mark"].upper() for row in page_rows}
            page_rows.extend(
                row
                for row in _glazing_rows_from_text(page_text, pdf_path.name)
                if row["mark"].upper() not in table_marks
            )
            rows.extend(page_rows)
    return _dedupe_rows(rows, GLAZING_HEADERS[:-1])


def _window_from_glazing(row: dict[str, str]) -> dict[str, str]:
    return {
        "mark": row["mark"],
        "quantity": row["quantity"],
        "width": row["width"],
        "height": row["height"],
        "type": row["description"],
        "material": row.get("material", ""),
        "glass_type": row.get("glass_type") or row.get("brand_product", ""),
        "finish": row.get("finish", ""),
        "noa": row["noa"],
        "brand_product": row["brand_product"],
        "level": row["level"],
        "remarks": row["remarks"],
        "extraction_status": "TABLE PARSED - glazing schedule",
        "source": row["source"],
    }


def _storefront_from_glazing(row: dict[str, str]) -> dict[str, str]:
    return _window_from_glazing(row)


def _door_from_glazing(row: dict[str, str]) -> dict[str, str]:
    door = _empty_door_row()
    door.update(
        {
            "door_no": row["mark"],
            "type": row["description"],
            "width": row["width"],
            "height": row["height"],
            "quantity": row["quantity"],
            "noa": row["noa"],
            "remarks": row["remarks"],
            "level": row["level"],
            "scope": "glazing schedule exterior opening",
            "extraction_status": "TABLE PARSED - glazing schedule",
            "raw_text": _join_notes(row["description"], row["brand_product"]),
            "source": row["source"],
        }
    )
    return door


def _extract_zone_tokens(texts: list[tuple[str, str]]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for source, text in texts:
        counts = Counter(re.findall(r"\b\d+[WD]\b", text))
        for zone_tag, count in sorted(counts.items()):
            rows.append(
                {
                    "source": source,
                    "zone_tag": zone_tag,
                    "occurrences": str(count),
                    "numeric_psf": "",
                    "status": "CRITICAL - numeric PSF value not shown in extracted text",
                }
            )
    return rows


def _extract_noas(text: str, source: str) -> list[dict[str, str]]:
    normalized = text.replace("−", "-")
    matches = re.finditer(
        r"NOA\s*#\s*([0-9-]+\.[0-9]+).*?EXPIRES\s+(\d{2}-\d{2}-\d{4})",
        normalized,
        flags=re.DOTALL,
    )
    rows: list[dict[str, str]] = []
    for match in matches:
        noa, expires_text = match.groups()
        expires = date_parser.parse(expires_text).date()
        context = normalized[max(0, match.start() - 300) : match.start()]
        if "ROOFING:" in context:
            scope = "roofing - outside opening quote scope"
        elif "POOL DECK:" in context:
            scope = "pool deck - outside opening quote scope"
        else:
            scope = "opening product"
        if expires < date.today() and scope == "opening product":
            status = "CRITICAL - expired opening product"
        elif expires < date.today():
            status = "REVIEW - expired outside opening quote scope"
        else:
            status = "OK"
        rows.append(
            {
                "noa": noa,
                "expires": expires.isoformat(),
                "scope": scope,
                "status": status,
                "source": source,
            }
        )
    return rows


def _load_sheet_index(job_dir: Path) -> list[dict[str, str]]:
    sheet_index_path = job_dir / "07_Organized_Plan_Set" / "sheet_index.csv"
    if not sheet_index_path.exists():
        raise IntakeError(f"Missing indexed plan-sheet inventory: {sheet_index_path}")
    with sheet_index_path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _unique_paths(paths: list[Path]) -> list[Path]:
    unique: list[Path] = []
    seen: set[Path] = set()
    for path in paths:
        if path not in seen:
            seen.add(path)
            unique.append(path)
    return unique


def _indexed_sheet_paths(
    job_dir: Path, sheet_index: list[dict[str, str]], compartment: str
) -> list[Path]:
    organized_dir = job_dir / "07_Organized_Plan_Set"
    matches: list[Path] = []
    for row in sheet_index:
        compartments = {
            value.strip()
            for value in row.get("compartments", "").split(",")
            if value.strip()
        }
        if compartment not in compartments:
            continue
        indexed_sheet_path = organized_dir / row["indexed_sheet"]
        if indexed_sheet_path.exists():
            matches.append(indexed_sheet_path)
    return _unique_paths(matches)


def _dedupe_rows(
    rows: list[dict[str, str]], identity_headers: tuple[str, ...]
) -> list[dict[str, str]]:
    unique_rows: list[dict[str, str]] = []
    seen: set[tuple[str, ...]] = set()
    for row in rows:
        identity = tuple(row.get(header, "") for header in identity_headers)
        if identity not in seen:
            seen.add(identity)
            unique_rows.append(row)
    return unique_rows


def _copy_extracted_pages(pdf_paths: list[Path], destination: Path) -> None:
    for pdf_path in _unique_paths(pdf_paths):
        copy_destination = destination / pdf_path.name
        counter = 2
        while copy_destination.exists():
            copy_destination = destination / f"{pdf_path.stem}_{counter}{pdf_path.suffix}"
            counter += 1
        copy2(pdf_path, copy_destination)


def _formula_map(workbook_path: Path) -> dict[tuple[str, str], str]:
    workbook = load_workbook(workbook_path, data_only=False)
    return {
        (sheet.title, cell.coordinate): cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    }


def _set_blank(sheet, coordinate: str, value: str) -> None:
    if value in (None, ""):
        return
    cell = sheet[coordinate]
    if cell.value not in (None, ""):
        raise IntakeError(
            f"Refusing to overwrite populated workbook cell {sheet.title}!{coordinate}."
        )
    cell.value = value


def _available_workbook_columns(sheet, populated_rows: tuple[int, ...]) -> list[str]:
    return [
        get_column_letter(column_index)
        for column_index in range(2, sheet.max_column + 1)
        if all(sheet.cell(row_index, column_index).value in (None, "") for row_index in populated_rows)
    ]


def _clone_blank_workbook_column(sheet, source_index: int, target_index: int) -> None:
    source_letter = get_column_letter(source_index)
    target_letter = get_column_letter(target_index)
    source_dimension = sheet.column_dimensions[source_letter]
    target_dimension = sheet.column_dimensions[target_letter]
    target_dimension.width = source_dimension.width
    target_dimension.hidden = source_dimension.hidden
    target_dimension.outlineLevel = source_dimension.outlineLevel
    target_dimension.bestFit = source_dimension.bestFit
    for row_index in range(1, sheet.max_row + 1):
        source_cell = sheet.cell(row_index, source_index)
        target_cell = sheet.cell(row_index, target_index)
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        if isinstance(source_cell.value, str) and source_cell.value.startswith("="):
            try:
                target_cell.value = Translator(
                    source_cell.value, origin=source_cell.coordinate
                ).translate_formula(target_cell.coordinate)
            except TranslatorError:
                target_cell.value = source_cell.value


def _extend_header_merge(sheet) -> None:
    for merged_range in list(sheet.merged_cells.ranges):
        if (
            merged_range.min_row == 1
            and merged_range.max_row == 1
            and merged_range.min_col == 1
            and merged_range.max_col < sheet.max_column
        ):
            sheet.unmerge_cells(str(merged_range))
            sheet.merge_cells(
                start_row=1,
                start_column=1,
                end_row=1,
                end_column=sheet.max_column,
            )
            return


def _ensure_workbook_columns(
    sheet, populated_rows: tuple[int, ...], required_count: int
) -> list[str]:
    available = _available_workbook_columns(sheet, populated_rows)
    while len(available) < required_count:
        source_index = sheet.max_column
        _clone_blank_workbook_column(sheet, source_index, source_index + 1)
        available = _available_workbook_columns(sheet, populated_rows)
    _extend_header_merge(sheet)
    return available


def _row_quantity(row: dict[str, str]) -> int:
    try:
        quantity = int(str(row.get("quantity") or "1").strip())
    except ValueError:
        return 1
    return max(quantity, 1)


def _write_count_check(
    sheet,
    title_row: int,
    identifier_row: int,
    rows: list[dict[str, str]],
    identifier_key: str,
) -> None:
    entries = [row for row in rows if row.get(identifier_key)]
    clear_end_row = max(sheet.max_row, title_row + len(entries) + 3)
    for row_index in range(title_row, clear_end_row + 1):
        for column_index in range(1, 4):
            sheet.cell(row_index, column_index).value = None
    sheet.cell(title_row, 1).value = "Count Check"
    header_row = title_row + 1
    sheet.cell(header_row, 1).value = "Mark / Unit ID"
    sheet.cell(header_row, 2).value = "Schedule Qty"
    sheet.cell(header_row, 3).value = "Workbook Columns"
    last_column = get_column_letter(sheet.max_column)
    first_data_row = title_row + 2
    for offset, row in enumerate(entries):
        row_index = first_data_row + offset
        sheet.cell(row_index, 1).value = row[identifier_key]
        sheet.cell(row_index, 2).value = _row_quantity(row)
        sheet.cell(row_index, 3).value = (
            f'=COUNTIF($B${identifier_row}:${last_column}${identifier_row},A{row_index})'
        )
    total_row = first_data_row + len(entries)
    if entries:
        sheet.cell(total_row, 1).value = "TOTAL"
        sheet.cell(total_row, 2).value = f"=SUM(B{first_data_row}:B{total_row - 1})"
        sheet.cell(total_row, 3).value = f"=SUM(C{first_data_row}:C{total_row - 1})"


def _write_door_count_check(sheet, doors: list[dict[str, str]]) -> None:
    door_types: list[str] = []
    seen: set[str] = set()
    for door in doors:
        door_type = str(door.get("type") or "").strip().upper()
        if not door_type or door_type in seen:
            continue
        seen.add(door_type)
        door_types.append(door_type)

    title_row = 20
    header_row = title_row + 1
    first_data_row = title_row + 2
    total_row = first_data_row + len(door_types)
    clear_end_row = max(sheet.max_row, total_row + 1)
    for row_index in range(title_row, clear_end_row + 1):
        for column_index in range(1, 3):
            sheet.cell(row_index, column_index).value = None

    sheet.cell(title_row, 1).value = "Count Check"
    sheet.cell(header_row, 1).value = "Type"
    sheet.cell(header_row, 2).value = "Count"
    last_column = get_column_letter(sheet.max_column)
    for offset, door_type in enumerate(door_types):
        row_index = first_data_row + offset
        sheet.cell(row_index, 1).value = door_type
        sheet.cell(row_index, 2).value = f"=COUNTIF($B$5:${last_column}$5,A{row_index})"
    if door_types:
        sheet.cell(total_row, 1).value = "Total : "
        sheet.cell(total_row, 2).value = f"=SUM(B{first_data_row}:B{total_row - 1})"


def _candidate_notes(row: dict[str, str]) -> str:
    return _join_notes(
        f"Quantity: {row['quantity']}" if row.get("quantity") else "",
        f"Level: {row['level']}" if row.get("level") else "",
        f"Product approval: {row['noa']}" if row.get("noa") else "",
        f"Brand/product: {row['brand_product']}" if row.get("brand_product") else "",
        row.get("remarks", ""),
    )


def _missing_source_field_note(
    row: dict[str, str], fields: tuple[tuple[str, str], ...]
) -> str:
    missing = [label for label, field in fields if not row.get(field)]
    if not missing:
        return ""
    return f"Missing explicit source fields: {', '.join(missing)}"


def _window_notes(row: dict[str, str]) -> str:
    return _join_notes(
        _candidate_notes(row),
        f"Panels: {row['panels']}" if row.get("panels") else "",
        _missing_source_field_note(
            row,
            (
                ("Material", "material"),
                ("Glass Type", "glass_type"),
                ("Frame Finish", "finish"),
            ),
        ),
        f"Source: {row['source']}" if row.get("source") else "",
    )


def _storefront_notes(row: dict[str, str]) -> str:
    return _join_notes(
        _candidate_notes(row),
        _missing_source_field_note(
            row,
            (
                ("Material", "material"),
                ("Glass Type", "glass_type"),
                ("Finish", "finish"),
            ),
        ),
        f"Source: {row['source']}" if row.get("source") else "",
    )


def _quote_door_rows(doors: list[dict[str, str]]) -> list[dict[str, str]]:
    quote_scopes = {"door schedule", "garage door schedule"}
    return [row for row in doors if row.get("scope") in quote_scopes]


def _door_requires_manual_parsing(row: dict[str, str]) -> bool:
    return row.get("extraction_status", "").upper().startswith("REVIEW")


def _workbook_capacity_issues(
    template_path: Path,
    windows: list[dict[str, str]],
    storefronts: list[dict[str, str]],
    doors: list[dict[str, str]],
) -> list[dict[str, str]]:
    workbook = load_workbook(template_path, data_only=False)
    groups = (
        (
            "window",
            windows,
            workbook["Windows"],
            (5, 6, 7, 8, 9, 10, 11, 12, 17),
            True,
        ),
        (
            "glazing / storefront",
            storefronts,
            workbook["Storefronts"],
            (4, 5, 6, 7, 8, 9, 10, 11, 12, 16),
            True,
        ),
        (
            "door schedule",
            _quote_door_rows(doors),
            workbook["Doors"],
            (4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 17),
            True,
        ),
    )
    issues: list[dict[str, str]] = []
    for label, rows, sheet, populated_rows, expandable in groups:
        columns = (
            _ensure_workbook_columns(sheet, populated_rows, len(rows))
            if expandable
            else _available_workbook_columns(sheet, populated_rows)
        )
        capacity = len(columns)
        if len(rows) <= capacity:
            continue
        issues.append(
            {
                "severity": "CRITICAL",
                "issue_type": "Quote workbook column capacity exceeded",
                "mark": f"{len(rows)} {label} candidates",
                "source_page": f"04_Quote_Workbook/{template_path.name}",
                "description": f"The quote template has {capacity} available {label} columns but {len(rows)} candidates were extracted.",
                "recommended_action": "Review the schedule CSV and extend the quote workbook template before issuing the quote.",
            }
        )
    return issues


def _identified_source_extraction_issues(
    door_pages: list[Path],
    window_pages: list[Path],
    storefront_pages: list[Path],
    pressure_pages: list[Path],
    doors: list[dict[str, str]],
    windows: list[dict[str, str]],
    storefronts: list[dict[str, str]],
    zones: list[dict[str, str]],
) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    schedule_groups = (
        ("door", door_pages, doors, "02_Schedules/doors_schedule_candidates.csv"),
        ("window", window_pages, windows, "02_Schedules/windows_schedule.csv"),
        (
            "storefront",
            storefront_pages,
            storefronts,
            "02_Schedules/storefront_schedule.csv",
        ),
    )
    for label, pages, rows, output_path in schedule_groups:
        if not pages or rows:
            continue
        issues.append(
            {
                "severity": "CRITICAL",
                "issue_type": "Recognized schedule requires manual extraction",
                "mark": f"{label} schedule",
                "source_page": ", ".join(path.name for path in pages),
                "description": f"A {label} schedule sheet was identified, but no {label} rows could be transported into {output_path} confidently.",
                "recommended_action": f"Review the indexed {label} schedule sheet and enter verified rows manually before pricing.",
            }
        )
    if pressure_pages and not zones:
        issues.append(
            {
                "severity": "CRITICAL",
                "issue_type": "Wind-pressure mapping requires manual review",
                "mark": "PSF / zone values",
                "source_page": ", ".join(path.name for path in pressure_pages),
                "description": "Dedicated wind-pressure sheets were identified, but numeric opening-to-zone PSF rows could not be transported into structured output confidently.",
                "recommended_action": "Review the indexed wind-pressure sheets and enter verified positive and negative PSF mappings manually before pricing.",
            }
        )
    return issues


def _structured_extraction_issues(
    schedule_pages: list[Path], structured_result: ExtractionResult
) -> list[dict[str, str]]:
    items = [item for schedule in structured_result.schedules for item in schedule.items]
    uncertain_items = [item for item in items if item.confidence < HUMAN_REVIEW_THRESHOLD]
    issues: list[dict[str, str]] = []
    if schedule_pages and not items:
        issues.append(
            {
                "severity": "CRITICAL",
                "issue_type": "Structured schedule extraction produced no rows",
                "mark": "Schedule extraction",
                "source_page": "02_Schedules/extraction_audit.json",
                "description": "Schedule pages were identified, but no rows passed through the structured extraction layer.",
                "recommended_action": "Review extraction_audit.json and enter verified rows manually before pricing.",
            }
        )
    if uncertain_items:
        issues.append(
            {
                "severity": "CRITICAL",
                "issue_type": "Low-confidence structured schedule rows require review",
                "mark": f"{len(uncertain_items)} rows",
                "source_page": "02_Schedules/uncertain_schedule_rows.csv",
                "description": f"{len(uncertain_items)} extracted schedule rows scored below {HUMAN_REVIEW_THRESHOLD:.2f} and were excluded from quote workbook population.",
                "recommended_action": "Review raw values and warnings, then correct or enter verified values manually before pricing.",
            }
        )
    return issues


def _missing_workbook_attribute_issues(
    windows: list[dict[str, str]],
    storefronts: list[dict[str, str]],
    doors: list[dict[str, str]],
) -> list[dict[str, str]]:
    groups = (
        (
            "Window",
            windows,
            (
                ("Material", "material"),
                ("Glass Type", "glass_type"),
                ("Frame Finish", "finish"),
            ),
            "Windows worksheet rows 10-12",
        ),
        (
            "Storefront",
            storefronts,
            (
                ("Material", "material"),
                ("Glass Type", "glass_type"),
                ("Finish", "finish"),
            ),
            "Storefronts worksheet rows 10-12",
        ),
        (
            "Door",
            _quote_door_rows(doors),
            (("Frame Finish Color", "frame_finish"),),
            "Doors worksheet row 13",
        ),
    )
    issues: list[dict[str, str]] = []
    for label, rows, fields, source_page in groups:
        missing_counts = {
            field_label: sum(1 for row in rows if not row.get(field_name))
            for field_label, field_name in fields
        }
        missing_counts = {
            field_label: count
            for field_label, count in missing_counts.items()
            if count
        }
        if not missing_counts:
            continue
        missing_summary = ", ".join(
            f"{field_label}: {count}" for field_label, count in missing_counts.items()
        )
        issues.append(
            {
                "severity": "HIGH",
                "issue_type": f"{label} finish/material review required",
                "mark": f"{len(rows)} {label.lower()} rows",
                "source_page": source_page,
                "description": f"Some {label.lower()} quote rows do not have explicit source values for required finish/material attributes ({missing_summary}).",
                "recommended_action": "Review the source schedule and template requirements; enter only verified material, glass type, finish, or frame-finish color values before pricing.",
            }
        )
    return issues


def _write_door_workbook_transfer_audit(
    draft_path: Path, doors: list[dict[str, str]]
) -> Path:
    workbook = load_workbook(draft_path, data_only=False)
    sheet = workbook["Doors"]
    workbook_columns = [
        get_column_letter(column_index)
        for column_index in range(2, sheet.max_column + 1)
        if sheet.cell(4, column_index).value not in (None, "")
    ]
    field_rows = {
        "door_no": 4,
        "type": 5,
        "location": 7,
        "width": 8,
        "height": 9,
        "thickness": 10,
        "door_material": 11,
        "frame_material": 12,
        "frame_finish": 13,
    }
    mismatches: list[dict[str, str]] = []
    audit_rows: list[dict[str, object]] = []
    if len(workbook_columns) != len(doors):
        mismatches.append(
            {
                "door_no": "ALL",
                "field": "row_count",
                "expected": str(len(doors)),
                "actual": str(len(workbook_columns)),
            }
        )
    for column, door in zip(workbook_columns, doors):
        row_mismatches: list[dict[str, str]] = []
        for field, row_index in field_rows.items():
            expected = str(door.get(field) or "")
            actual = str(sheet[f"{column}{row_index}"].value or "")
            if actual != expected:
                mismatch = {
                    "door_no": str(door["door_no"]),
                    "field": field,
                    "expected": expected,
                    "actual": actual,
                }
                mismatches.append(mismatch)
                row_mismatches.append(mismatch)
        audit_rows.append(
            {
                "door_no": door["door_no"],
                "workbook_column": column,
                "matched": not row_mismatches,
                "mismatches": row_mismatches,
            }
        )
    audit = {
        "passed": not mismatches,
        "expected_door_count": len(doors),
        "workbook_door_count": len(workbook_columns),
        "expected_marks": [row["door_no"] for row in doors],
        "workbook_marks": [str(sheet[f"{column}4"].value) for column in workbook_columns],
        "checked_fields": list(field_rows),
        "mismatches": mismatches,
        "rows": audit_rows,
    }
    audit_path = (
        draft_path.parents[1] / "02_Schedules" / "door_workbook_transfer_audit.json"
    )
    audit_path.write_text(json.dumps(audit, indent=2) + "\n", encoding="utf-8")
    if mismatches:
        raise IntakeError(
            "Door workbook transfer reconciliation failed. "
            f"Review {audit_path} before issuing the quote."
        )
    return audit_path


def _write_window_workbook_transfer_audit(
    draft_path: Path,
    windows: list[dict[str, str]],
    glazing_requirements: dict[str, str],
) -> Path:
    workbook = load_workbook(draft_path, data_only=False)
    sheet = workbook["Windows"]
    workbook_columns = [
        get_column_letter(column_index)
        for column_index in range(2, sheet.max_column + 1)
        if sheet.cell(6, column_index).value not in (None, "")
    ]
    field_rows = {
        "type": 5,
        "mark": 6,
        "source_schedule": 7,
        "width": 8,
        "height": 9,
        "material": 10,
        "glass_type": 11,
        "finish": 12,
    }
    mismatches: list[dict[str, str]] = []
    audit_rows: list[dict[str, object]] = []
    if len(workbook_columns) != len(windows):
        mismatches.append(
            {
                "mark": "ALL",
                "field": "row_count",
                "expected": str(len(windows)),
                "actual": str(len(workbook_columns)),
            }
        )
    for column, window in zip(workbook_columns, windows):
        row_mismatches: list[dict[str, str]] = []
        for field, row_index in field_rows.items():
            expected = str(window.get(field) or "")
            actual = str(sheet[f"{column}{row_index}"].value or "")
            if actual != expected:
                mismatch = {
                    "mark": str(window["mark"]),
                    "field": field,
                    "expected": expected,
                    "actual": actual,
                }
                mismatches.append(mismatch)
                row_mismatches.append(mismatch)
        expected_notes = _window_notes(window)
        actual_notes = str(sheet[f"{column}17"].value or "")
        if actual_notes != expected_notes:
            mismatch = {
                "mark": str(window["mark"]),
                "field": "notes",
                "expected": expected_notes,
                "actual": actual_notes,
            }
            mismatches.append(mismatch)
            row_mismatches.append(mismatch)
        for field, row_index in (("u_factor", 14), ("shgc", 15)):
            expected = str(glazing_requirements.get(field) or "")
            actual = str(sheet[f"{column}{row_index}"].value or "")
            if actual != expected:
                mismatch = {
                    "mark": str(window["mark"]),
                    "field": field,
                    "expected": expected,
                    "actual": actual,
                }
                mismatches.append(mismatch)
                row_mismatches.append(mismatch)
        audit_rows.append(
            {
                "mark": window["mark"],
                "workbook_column": column,
                "matched": not row_mismatches,
                "mismatches": row_mismatches,
            }
        )
    audit = {
        "passed": not mismatches,
        "expected_window_count": len(windows),
        "workbook_window_count": len(workbook_columns),
        "expected_marks": [row["mark"] for row in windows],
        "workbook_marks": [str(sheet[f"{column}6"].value) for column in workbook_columns],
        "checked_fields": [*field_rows, "notes", "u_factor", "shgc"],
        "mismatches": mismatches,
        "rows": audit_rows,
    }
    audit_path = (
        draft_path.parents[1] / "02_Schedules" / "window_workbook_transfer_audit.json"
    )
    audit_path.write_text(json.dumps(audit, indent=2) + "\n", encoding="utf-8")
    if mismatches:
        raise IntakeError(
            "Window workbook transfer reconciliation failed. "
            f"Review {audit_path} before issuing the quote."
        )
    return audit_path


def _write_storefront_workbook_transfer_audit(
    draft_path: Path,
    storefronts: list[dict[str, str]],
    glazing_requirements: dict[str, str],
) -> Path:
    workbook = load_workbook(draft_path, data_only=False)
    sheet = workbook["Storefronts"]
    workbook_columns = [
        get_column_letter(column_index)
        for column_index in range(2, sheet.max_column + 1)
        if sheet.cell(4, column_index).value not in (None, "")
    ]
    field_rows = {
        "mark": 4,
        "type": 5,
        "source_schedule": 6,
        "panels": 7,
        "width": 8,
        "height": 9,
        "material": 10,
        "glass_type": 11,
        "finish": 12,
    }
    mismatches: list[dict[str, str]] = []
    audit_rows: list[dict[str, object]] = []
    if len(workbook_columns) != len(storefronts):
        mismatches.append(
            {
                "mark": "ALL",
                "field": "row_count",
                "expected": str(len(storefronts)),
                "actual": str(len(workbook_columns)),
            }
        )
    for column, storefront in zip(workbook_columns, storefronts):
        row_mismatches: list[dict[str, str]] = []
        for field, row_index in field_rows.items():
            expected = str(storefront.get(field) or "")
            actual = str(sheet[f"{column}{row_index}"].value or "")
            if actual != expected:
                mismatch = {
                    "mark": str(storefront["mark"]),
                    "field": field,
                    "expected": expected,
                    "actual": actual,
                }
                mismatches.append(mismatch)
                row_mismatches.append(mismatch)
        expected_notes = _storefront_notes(storefront)
        actual_notes = str(sheet[f"{column}16"].value or "")
        if actual_notes != expected_notes:
            mismatch = {
                "mark": str(storefront["mark"]),
                "field": "notes",
                "expected": expected_notes,
                "actual": actual_notes,
            }
            mismatches.append(mismatch)
            row_mismatches.append(mismatch)
        for field, row_index in (("u_factor", 13), ("shgc", 14)):
            expected = str(glazing_requirements.get(field) or "")
            actual = str(sheet[f"{column}{row_index}"].value or "")
            if actual != expected:
                mismatch = {
                    "mark": str(storefront["mark"]),
                    "field": field,
                    "expected": expected,
                    "actual": actual,
                }
                mismatches.append(mismatch)
                row_mismatches.append(mismatch)
        audit_rows.append(
            {
                "mark": storefront["mark"],
                "workbook_column": column,
                "matched": not row_mismatches,
                "mismatches": row_mismatches,
            }
        )
    audit = {
        "passed": not mismatches,
        "expected_storefront_package_count": len(storefronts),
        "workbook_storefront_package_count": len(workbook_columns),
        "expected_marks": [row["mark"] for row in storefronts],
        "workbook_marks": [str(sheet[f"{column}4"].value) for column in workbook_columns],
        "glazing_requirements": glazing_requirements,
        "checked_fields": [*field_rows, "notes", "u_factor", "shgc"],
        "mismatches": mismatches,
        "rows": audit_rows,
    }
    audit_path = (
        draft_path.parents[1] / "02_Schedules" / "storefront_workbook_transfer_audit.json"
    )
    audit_path.write_text(json.dumps(audit, indent=2) + "\n", encoding="utf-8")
    if mismatches:
        raise IntakeError(
            "Storefront workbook transfer reconciliation failed. "
            f"Review {audit_path} before issuing the quote."
        )
    return audit_path


def _create_workbook_draft(
    template_path: Path,
    draft_path: Path,
    windows: list[dict[str, str]],
    storefronts: list[dict[str, str]],
    doors: list[dict[str, str]],
    critical_issues: list[str],
    glazing_requirements: dict[str, str] | None = None,
) -> None:
    glazing_requirements = glazing_requirements or {}
    before_formulas = _formula_map(template_path)
    workbook = load_workbook(template_path, data_only=False)

    windows_sheet = workbook["Windows"]
    window_columns = _ensure_workbook_columns(
        windows_sheet, (5, 6, 7, 8, 9, 10, 11, 12, 17), len(windows)
    )
    for column, row in zip(window_columns, windows):
        _set_blank(windows_sheet, f"{column}5", row["type"])
        _set_blank(windows_sheet, f"{column}6", row["mark"])
        _set_blank(windows_sheet, f"{column}7", row.get("source_schedule", ""))
        _set_blank(windows_sheet, f"{column}8", row["width"])
        _set_blank(windows_sheet, f"{column}9", row["height"])
        _set_blank(windows_sheet, f"{column}10", row["material"])
        _set_blank(windows_sheet, f"{column}11", row.get("glass_type", ""))
        _set_blank(windows_sheet, f"{column}12", row.get("finish", ""))
        _set_blank(windows_sheet, f"{column}14", glazing_requirements.get("u_factor", ""))
        _set_blank(windows_sheet, f"{column}15", glazing_requirements.get("shgc", ""))
        _set_blank(windows_sheet, f"{column}17", _window_notes(row))
    _write_count_check(windows_sheet, 20, 6, windows, "mark")

    storefronts_sheet = workbook["Storefronts"]
    storefronts_sheet["A1"] = "STOREFRONT GLAZING PACKAGE - UNIT COLUMNS"
    storefront_columns = _ensure_workbook_columns(
        storefronts_sheet, (4, 5, 6, 7, 8, 9, 10, 11, 12, 16), len(storefronts)
    )
    for column, row in zip(storefront_columns, storefronts):
        _set_blank(storefronts_sheet, f"{column}4", row["mark"])
        _set_blank(storefronts_sheet, f"{column}5", row["type"])
        _set_blank(storefronts_sheet, f"{column}6", row.get("source_schedule", ""))
        _set_blank(storefronts_sheet, f"{column}7", row.get("panels", ""))
        _set_blank(storefronts_sheet, f"{column}8", row["width"])
        _set_blank(storefronts_sheet, f"{column}9", row["height"])
        _set_blank(storefronts_sheet, f"{column}10", row["material"])
        _set_blank(storefronts_sheet, f"{column}11", row.get("glass_type", ""))
        _set_blank(storefronts_sheet, f"{column}12", row.get("finish", ""))
        _set_blank(storefronts_sheet, f"{column}13", glazing_requirements.get("u_factor", ""))
        _set_blank(storefronts_sheet, f"{column}14", glazing_requirements.get("shgc", ""))
        _set_blank(storefronts_sheet, f"{column}16", _storefront_notes(row))
    _write_count_check(storefronts_sheet, 19, 4, storefronts, "mark")

    doors_sheet = workbook["Doors"]
    doors_sheet["A1"] = "DOOR SCHEDULE - UNIT COLUMNS"
    quote_doors = _quote_door_rows(doors)
    door_columns = _ensure_workbook_columns(
        doors_sheet,
        (4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 17),
        len(quote_doors),
    )
    for column, row in zip(door_columns, quote_doors):
        _set_blank(doors_sheet, f"{column}4", row["door_no"])
        _set_blank(doors_sheet, f"{column}5", row["type"])
        sliding_value = "YES" if "SLIDING" in row["type"].upper() else ""
        _set_blank(doors_sheet, f"{column}6", sliding_value)
        _set_blank(doors_sheet, f"{column}7", row["location"])
        _set_blank(doors_sheet, f"{column}8", row["width"])
        _set_blank(doors_sheet, f"{column}9", row["height"])
        _set_blank(doors_sheet, f"{column}10", row["thickness"])
        _set_blank(doors_sheet, f"{column}11", row["door_material"])
        _set_blank(doors_sheet, f"{column}12", row["frame_material"])
        _set_blank(doors_sheet, f"{column}13", row["frame_finish"])
        door_notes = _join_notes(
            _candidate_notes(row),
            f"Panels: {row['panels']}" if row.get("panels") else "",
            f"Fixed panels: {row['fixed_panels']}" if row.get("fixed_panels") else "",
            f"Hardware: {row['panic_hardware']}" if row.get("panic_hardware") else "",
            f"Details: {row['raw_text']}" if row.get("raw_text") else "",
            "Missing explicit source field: Frame Finish Color"
            if not row.get("frame_finish")
            else "",
            f"Source: {row['source']}",
        )
        _set_blank(doors_sheet, f"{column}17", door_notes)
    _write_door_count_check(doors_sheet, quote_doors)

    if "Automation Review" in workbook.sheetnames:
        del workbook["Automation Review"]
    review = workbook.create_sheet("Automation Review")
    review.append(["HurricaneOps Quote Workbook Draft"])
    review.append(["Generated", datetime.now().isoformat(timespec="seconds")])
    review.append(["Status", "REVIEW REQUIRED - pricing and quantities are not complete"])
    review.append([])
    review.append(["CRITICAL issues"])
    for issue in critical_issues:
        review.append([issue])
    review.append([])
    review.append(["Notes"])
    review.append(["Only fields explicitly present in extracted drawing tables were filled."])
    review.append(["Window and storefront Glass TYPE is populated from the schedule Brand/Product column when no separate glass-type column is present."])
    review.append(["Window/storefront Material and Finish remain blank unless explicitly extracted; missing values are listed in QA and item notes."])
    review.append(["Door FRAME FINISH is populated from the second value in combined DOOR / FRAME material cells when available."])
    review.append(["The project-specific storefront glazing package was kept together in drawing order on the Storefronts worksheet."])
    review.append(["Window-only glazing candidates were mapped to the Windows worksheet and reconciled after save."])
    review.append(["Architectural and garage-door schedule candidates were mapped in drawing order without overwriting populated cells."])
    review.append(["PSF and zone assignments remain blank for manual review."])
    review.append(["Existing formulas were preserved."])
    review.column_dimensions["A"].width = 110

    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(draft_path)
    after_formulas = _formula_map(draft_path)
    if before_formulas != {
        key: value for key, value in after_formulas.items() if key in before_formulas
    }:
        raise IntakeError("Formula verification failed after workbook draft creation.")
    _write_window_workbook_transfer_audit(draft_path, windows, glazing_requirements)
    _write_storefront_workbook_transfer_audit(
        draft_path, storefronts, glazing_requirements
    )
    _write_door_workbook_transfer_audit(draft_path, quote_doors)


def _write_reports(
    job_dir: Path,
    project_name: str,
    address: str,
    windows: list[dict[str, str]],
    storefronts: list[dict[str, str]],
    doors: list[dict[str, str]],
    glazing_rows: list[dict[str, str]],
    zones: list[dict[str, str]],
    noas: list[dict[str, str]],
    missing_sources: list[str],
    extra_qa_issues: list[dict[str, str]],
) -> tuple[list[str], list[dict[str, str]]]:
    qa_issues = [
        {
            "severity": "CRITICAL",
            "issue_type": "Missing indexed plan sheet",
            "mark": source_label,
            "source_page": "07_Organized_Plan_Set/sheet_index.csv",
            "description": f"No {source_label} sheet was identified from the plan-set index, title blocks, or page text.",
            "recommended_action": "Review sheet_index.csv and identify the correct source page before pricing.",
        }
        for source_label in missing_sources
    ]
    qa_issues.extend(
        extra_qa_issues
    )
    qa_issues.extend(
        [
            {
                "severity": "CRITICAL",
                "issue_type": "Missing numeric PSF mapping",
                "mark": row["zone_tag"],
                "source_page": row["source"],
                "description": f"Pressure tag {row['zone_tag']} has no numeric PSF mapping.",
                "recommended_action": "Confirm positive and negative design pressures before pricing.",
            }
            for row in zones
        ]
    )
    qa_issues.extend(
        {
            "severity": "CRITICAL",
            "issue_type": "Expired opening-product NOA",
            "mark": row["noa"],
            "source_page": row["source"],
            "description": f"Opening-product NOA expired on {row['expires']}.",
            "recommended_action": "Select and document a current approved product before issue.",
        }
        for row in noas
        if row["status"].startswith("CRITICAL")
    )
    raw_door_count = sum(_door_requires_manual_parsing(row) for row in doors)
    if raw_door_count:
        qa_issues.append(
            {
                "severity": "CRITICAL",
                "issue_type": "Manual door schedule parsing required",
                "mark": f"{raw_door_count} rows",
                "source_page": "Indexed door schedule sheets",
                "description": f"{raw_door_count} common-area door rows were retained as raw text only.",
                "recommended_action": "Review raw candidate rows and confirm attributes before pricing.",
            }
        )
    qa_issues.extend(
        [
            {
                "severity": "HIGH",
                "issue_type": "Quantity verification required",
                "mark": "All openings",
                "source_page": "Indexed schedule sheets",
                "description": "Opening quantities were not guessed or auto-filled.",
                "recommended_action": "Confirm counts against plans and elevations.",
            },
            {
                "severity": "MEDIUM",
                "issue_type": "Door workbook mapping review required",
                "mark": "Architectural and garage doors",
                "source_page": "Indexed door schedule sheets",
                "description": "Validated architectural and garage-door schedule candidates were mapped into quote columns in drawing order when source values were explicit.",
                "recommended_action": "Review size, swing, PSF, zone, and notes before pricing.",
            },
        ]
    )
    qa_issues.extend(
        {
            "severity": "LOW",
            "issue_type": "Expired NOA outside opening quote scope",
            "mark": row["noa"],
            "source_page": row["source"],
            "description": f"Outside-scope NOA expired on {row['expires']}.",
            "recommended_action": "Flag for the relevant trade if this scope is included later.",
        }
        for row in noas
        if row["status"].startswith("REVIEW")
    )
    critical_issues = [
        issue["description"] + f" Source: {issue['source_page']}."
        for issue in qa_issues
        if issue["severity"] == "CRITICAL"
    ]

    schedules_dir = job_dir / "02_Schedules"
    qa_dir = job_dir / "03_Takeoff_QA"
    noa_dir = job_dir / "05_NOA_PSF_Check"
    proposal_dir = job_dir / "06_Proposal_Email"

    _write_csv(schedules_dir / "windows_schedule.csv", WINDOW_HEADERS, windows)
    _write_csv(schedules_dir / "storefront_schedule.csv", STOREFRONT_HEADERS, storefronts)
    _write_csv(schedules_dir / "doors_schedule_candidates.csv", DOOR_HEADERS, doors)
    _write_csv(
        schedules_dir / "glazing_schedule_candidates.csv", GLAZING_HEADERS, glazing_rows
    )
    _write_csv(
        noa_dir / "psf_zone_conflicts.csv",
        ("source", "zone_tag", "occurrences", "numeric_psf", "status"),
        zones,
    )
    _write_csv(
        noa_dir / "noa_expiration_check.csv",
        ("noa", "expires", "scope", "status", "source"),
        noas,
    )
    _write_xlsx_report(qa_dir / "qa_report.xlsx", "QA Issues", ISSUE_HEADERS, qa_issues)
    _write_xlsx_report(
        noa_dir / "psf_zone_report.xlsx",
        "PSF Zone Review",
        ("source", "zone_tag", "occurrences", "numeric_psf", "status"),
        zones,
    )

    qa_lines = [
        "# HurricaneOps QA Report",
        "",
        "## Status",
        f"- CRITICAL issues: {len(critical_issues)}",
        "- Quote status: REVIEW REQUIRED",
        "",
        "## Extracted Schedule Candidates",
        f"- Windows: {len(windows)}",
        f"- Storefronts: {len(storefronts)}",
        f"- Door candidates: {len(doors)}",
        f"- Glazing schedule rows: {len(glazing_rows)}",
        f"- Door rows requiring manual parsing: {raw_door_count}",
        "",
        "## CRITICAL Issues",
    ]
    qa_lines.extend(f"- {issue}" for issue in critical_issues)
    qa_lines.extend(
        [
            "",
            "## Manual Review Notes",
            "- Numeric PSF values and opening-to-zone assignments were not present in the extracted text.",
            "- Window and storefront quantities were not guessed.",
            "- Window and storefront material, glass type, and finish must be reviewed against the source schedule.",
            "- Door frame-finish color must be reviewed before pricing.",
            "- Door candidates remain separate by door number, dimensions, material, rating, and notes.",
            "- Review the original schedule sheets before pricing or issuing a proposal.",
            "",
        ]
    )
    (qa_dir / "qa_report.md").write_text("\n".join(qa_lines), encoding="utf-8")

    conflict_lines = [
        "# PSF / Zone Conflict Report",
        "",
        "## Status",
        "- CRITICAL: zone tags were found, but numeric PSF values were not available in extracted text.",
        "- No PSF values were guessed or assigned to quote workbook rows.",
        "",
        "## Required Review",
        "- Confirm numeric positive and negative pressures for each zone tag.",
        "- Confirm each opening mark's elevation and zone before combining quantities.",
        "- Keep records separate when PSF, zone, size, swing, or notes differ.",
        "",
    ]
    (noa_dir / "psf_zone_conflict_report.md").write_text(
        "\n".join(conflict_lines), encoding="utf-8"
    )

    proposal_label = " - ".join(
        value.strip() for value in (project_name, address) if value.strip()
    ) or job_dir.name
    proposal_lines = [
        f"Subject: Draft Proposal Review Required - {proposal_label}",
        "",
        "Hello,",
        "",
        "The initial HurricaneOps takeoff draft has been prepared for review.",
        "This is not a final proposal and should not be issued for pricing yet.",
        "",
        "Required review items:",
        "- Confirm window, storefront, and exterior door quantities.",
        "- Confirm window/storefront material, glass type, finish, and door frame-finish color.",
        "- Confirm pressure-zone assignments and numeric PSF values.",
        "- Review expired NOA entries and select current approved products.",
        "- Review schedule CSVs against the indexed original pages listed in sheet_index.csv.",
        "",
        "No missing values were guessed. Records with differing attributes must remain separate.",
        "",
        "Regards,",
        "",
    ]
    (proposal_dir / "proposal_email_draft.txt").write_text(
        "\n".join(proposal_lines), encoding="utf-8"
    )
    return critical_issues, qa_issues


def _write_completed_checklist(job_dir: Path, critical_count: int) -> None:
    checklist_path = job_dir / "checklist.md"
    checklist_path.write_text(
        "\n".join(
            (
                "# HurricaneOps Job Checklist",
                "",
                "- [x] PDF plan set copied",
                "- [x] Plan PDFs compartmentalized for estimator review",
                "- [x] Quote workbook template copied",
                "- [x] Relevant PDF schedule and pressure pages copied",
                "- [x] Text-backed schedule candidates extracted",
                "- [x] Schedule CSV files created",
                "- [x] Takeoff QA report created",
                "- [x] Formula-preserving quote workbook draft created",
                "- [x] NOA / PSF conflict reports created",
                "- [x] Proposal email draft created",
                f"- [ ] Resolve {critical_count} CRITICAL review items before pricing or issue",
                "",
            )
        ),
        encoding="utf-8",
    )


def _write_pdf_zip(zip_path: Path, pdf_paths: list[Path], relative_to: Path) -> None:
    with ZipFile(zip_path, "w", ZIP_DEFLATED) as archive:
        for pdf_path in sorted(pdf_paths):
            archive.write(pdf_path, pdf_path.relative_to(relative_to))


def _write_organized_plan_downloads(job_dir: Path) -> None:
    organized_dir = job_dir / "07_Organized_Plan_Set"
    download_dir = organized_dir / "00_Download_Packages"
    download_dir.mkdir(exist_ok=True)
    for key, folder_name, _label in ORGANIZED_PLAN_COMPARTMENTS:
        folder = organized_dir / folder_name
        _write_pdf_zip(
            download_dir / f"{key}.zip",
            list(folder.glob("*.pdf")),
            organized_dir,
        )
    organized_pdfs = [
        pdf_path
        for _key, folder_name, _label in ORGANIZED_PLAN_COMPARTMENTS
        for pdf_path in (organized_dir / folder_name).glob("*.pdf")
    ]
    _write_pdf_zip(
        download_dir / "organized_plan_set_pdfs.zip",
        organized_pdfs,
        organized_dir,
    )
    print(f"[HurricaneOps] Created organized PDF download packages: {download_dir}")


def _organized_source_counts(job_dir: Path) -> dict[str, int]:
    organized_dir = job_dir / "07_Organized_Plan_Set"
    return {
        key: len(list((organized_dir / folder_name).glob("*.pdf")))
        for key, folder_name, _label in ORGANIZED_PLAN_COMPARTMENTS
    }


def _organized_review_files() -> dict[str, str]:
    return {
        "floor_plans.zip": "07_Organized_Plan_Set/00_Download_Packages/floor_plans.zip",
        "door_schedules.zip": "07_Organized_Plan_Set/00_Download_Packages/door_schedules.zip",
        "window_schedules.zip": "07_Organized_Plan_Set/00_Download_Packages/window_schedules.zip",
        "storefront_schedules.zip": "07_Organized_Plan_Set/00_Download_Packages/storefront_schedules.zip",
        "architectural_elevations.zip": "07_Organized_Plan_Set/00_Download_Packages/architectural_elevations.zip",
        "wind_pressure_elevations.zip": "07_Organized_Plan_Set/00_Download_Packages/wind_pressure_elevations.zip",
        "other_plan_documents.zip": "07_Organized_Plan_Set/00_Download_Packages/other_plan_documents.zip",
        "organized_plan_set_pdfs.zip": "07_Organized_Plan_Set/00_Download_Packages/organized_plan_set_pdfs.zip",
        "source_file_manifest.csv": "07_Organized_Plan_Set/source_file_manifest.csv",
        "sheet_index.csv": "07_Organized_Plan_Set/sheet_index.csv",
    }


def _write_pipeline_summary(
    job_dir: Path,
    windows: list[dict[str, str]],
    storefronts: list[dict[str, str]],
    doors: list[dict[str, str]],
    quote_rows: dict[str, list[dict[str, str]]],
    glazing_rows: list[dict[str, str]],
    zones: list[dict[str, str]],
    noas: list[dict[str, str]],
    critical_issues: list[str],
    qa_issues: list[dict[str, str]],
    structured_row_count: int,
    structured_uncertain_count: int,
    project_name: str,
    address: str,
    notes: str,
    municipality_lookup: dict[str, object],
) -> Path:
    raw_door_count = sum(_door_requires_manual_parsing(row) for row in doors)
    organized_source_counts = _organized_source_counts(job_dir)
    issue_counts = {
        severity: sum(issue["severity"] == severity for issue in qa_issues)
        for severity in ("LOW", "MEDIUM", "HIGH", "CRITICAL")
    }
    summary_path = job_dir / "pipeline_summary.json"
    summary = {
        "job_id": job_dir.name,
        "job_name": job_dir.name,
        "job_dir": str(job_dir.resolve()),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "project_name": project_name or "Untitled Project",
        "address": address or "Address Not Provided",
        "municipality_location": municipality_lookup.get("display_name", ""),
        "municipality_lookup": municipality_lookup,
        "notes": notes,
        "organized_plan_schema_version": ORGANIZED_PLAN_SCHEMA_VERSION,
        "status": "REVIEW REQUIRED" if critical_issues else "READY FOR REVIEW",
        "critical_count": len(critical_issues),
        "critical_issues": critical_issues,
        "qa_issue_counts": issue_counts,
        "qa_issues": qa_issues,
        "counts": {
            "windows": len(windows),
            "storefronts": len(storefronts),
            "common_area_doors": len(doors),
            "quote_window_rows": len(quote_rows["windows"]),
            "quote_glazing_storefront_rows": len(quote_rows["storefronts"]),
            "quote_door_rows": len(_quote_door_rows(quote_rows["doors"])),
            "quote_exterior_doors": sum(
                row.get("scope") == "garage door schedule"
                for row in quote_rows["doors"]
            ),
            "doors_requiring_manual_parsing": raw_door_count,
            "glazing_schedule_rows": len(glazing_rows),
            "structured_schedule_rows": structured_row_count,
            "structured_uncertain_rows": structured_uncertain_count,
            "pressure_zone_rows": len(zones),
            "noa_rows": len(noas),
            "organized_source_pdfs": organized_source_counts,
            "indexed_plan_sheets": len(_load_sheet_index(job_dir)),
        },
        "review_files": {
            "filled_quote_workbook.xlsx": "04_Quote_Workbook/filled_quote_workbook.xlsx",
            "qa_report.xlsx": "03_Takeoff_QA/qa_report.xlsx",
            "psf_zone_report.xlsx": "05_NOA_PSF_Check/psf_zone_report.xlsx",
            "windows_schedule.csv": "02_Schedules/windows_schedule.csv",
            "storefront_schedule.csv": "02_Schedules/storefront_schedule.csv",
            "doors_schedule_candidates.csv": "02_Schedules/doors_schedule_candidates.csv",
            "glazing_schedule_candidates.csv": "02_Schedules/glazing_schedule_candidates.csv",
            "extracted_schedules.json": "02_Schedules/extracted_schedules.json",
            "extraction_audit.json": "02_Schedules/extraction_audit.json",
            "uncertain_schedule_rows.csv": "02_Schedules/uncertain_schedule_rows.csv",
            "window_workbook_transfer_audit.json": "02_Schedules/window_workbook_transfer_audit.json",
            "storefront_workbook_transfer_audit.json": "02_Schedules/storefront_workbook_transfer_audit.json",
            "glazing_thermal_requirements.json": "02_Schedules/glazing_thermal_requirements.json",
            "door_workbook_transfer_audit.json": "02_Schedules/door_workbook_transfer_audit.json",
            **_organized_review_files(),
            "municipality_lookup.json": "municipality_lookup.json",
            "proposal_email_draft.txt": "06_Proposal_Email/proposal_email_draft.txt",
            "full_job_package.zip": "full_job_package.zip",
        },
    }
    summary_path.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    return summary_path


def _write_job_zip(job_dir: Path) -> Path:
    zip_path = job_dir / "full_job_package.zip"
    with ZipFile(zip_path, "w", ZIP_DEFLATED) as archive:
        for path in sorted(job_dir.rglob("*")):
            if path.is_file() and path != zip_path:
                archive.write(path, path.relative_to(job_dir))
    return zip_path


def backfill_organized_plan_downloads(job_dir: Path) -> None:
    source_dir = job_dir / "00_Source_PDFs"
    plan_sets = sorted(
        path
        for path in source_dir.iterdir()
        if path.is_dir() and any(path.rglob("*.pdf"))
    )
    if len(plan_sets) != 1:
        raise IntakeError(
            f"Expected one preserved PDF plan set in {source_dir}, but found {len(plan_sets)}."
        )
    organized_dir = job_dir / "07_Organized_Plan_Set"
    if organized_dir.exists():
        backup_dir = job_dir / "07_Organized_Plan_Set_Previous"
        counter = 2
        while backup_dir.exists():
            backup_dir = job_dir / f"07_Organized_Plan_Set_Previous_{counter}"
            counter += 1
        move(str(organized_dir), str(backup_dir))
    organized_dir.mkdir()
    pdf_paths = sorted(plan_sets[0].rglob("*.pdf"))
    _organize_source_pdfs(job_dir, plan_sets[0], pdf_paths)
    _write_organized_plan_downloads(job_dir)

    summary_path = job_dir / "pipeline_summary.json"
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    summary["organized_plan_schema_version"] = ORGANIZED_PLAN_SCHEMA_VERSION
    summary["counts"]["organized_source_pdfs"] = _organized_source_counts(job_dir)
    summary["review_files"].update(_organized_review_files())
    summary_path.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    _write_job_zip(job_dir)


def _append_jobs_index(job_dir: Path, summary: dict[str, object]) -> None:
    headers = (
        "job_id",
        "project_name",
        "address",
        "municipality_location",
        "municipality_status",
        "created_at",
        "job_folder",
        "critical_count",
        "high_count",
        "filled_workbook_path",
        "qa_report_path",
    )
    existing_rows: list[dict[str, str]] = []
    if JOBS_INDEX_PATH.exists():
        with JOBS_INDEX_PATH.open(newline="", encoding="utf-8") as handle:
            existing_rows = list(csv.DictReader(handle))
    row = {
        "job_id": str(summary["job_id"]),
        "project_name": str(summary["project_name"]),
        "address": str(summary["address"]),
        "municipality_location": str(summary.get("municipality_location", "")),
        "municipality_status": str(
            summary.get("municipality_lookup", {}).get("status", "")
            if isinstance(summary.get("municipality_lookup"), dict)
            else ""
        ),
        "created_at": str(summary["generated_at"]),
        "job_folder": str(job_dir.resolve()),
        "critical_count": str(summary["critical_count"]),
        "high_count": str(summary["qa_issue_counts"]["HIGH"]),
        "filled_workbook_path": str((job_dir / "04_Quote_Workbook/filled_quote_workbook.xlsx").resolve()),
        "qa_report_path": str((job_dir / "03_Takeoff_QA/qa_report.xlsx").resolve()),
    }
    existing_rows = [item for item in existing_rows if item.get("job_id") != row["job_id"]]
    existing_rows.append(row)
    with JOBS_INDEX_PATH.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(existing_rows)


def run_bid_pipeline(
    inbox_dir: Path = INBOX_DIR,
    jobs_dir: Path = JOBS_DIR,
    project_name: str = "",
    address: str = "",
    notes: str = "",
) -> tuple[Path, list[str]]:
    job_dir, _plan_set_dir, template_path = create_job_from_inbox(
        inbox_dir=inbox_dir,
        jobs_dir=jobs_dir,
        project_name=project_name,
        address=address,
    )
    municipality_result = lookup_municipality(address).to_dict()
    (job_dir / "municipality_lookup.json").write_text(
        json.dumps(municipality_result, indent=2) + "\n",
        encoding="utf-8",
    )
    if municipality_result.get("display_name"):
        print(f"[HurricaneOps] Municipality: {municipality_result['display_name']}")
    elif municipality_result.get("warning"):
        print(f"[HurricaneOps] Municipality lookup: {municipality_result['warning']}")
    sheet_index = _load_sheet_index(job_dir)
    door_pages = _indexed_sheet_paths(job_dir, sheet_index, "door_schedules")
    window_pages = _indexed_sheet_paths(job_dir, sheet_index, "window_schedules")
    storefront_pages = _indexed_sheet_paths(job_dir, sheet_index, "storefront_schedules")
    pressure_pages = _indexed_sheet_paths(job_dir, sheet_index, "wind_pressure_elevations")
    schedule_pages = _unique_paths([*window_pages, *storefront_pages, *door_pages])
    _copy_extracted_pages(
        [*door_pages, *schedule_pages, *pressure_pages],
        job_dir / "01_Extracted_Pages",
    )
    structured_result = ScheduleExtractor(project_address=address).extract_files(
        schedule_pages
    )
    write_extraction_outputs(structured_result, job_dir / "02_Schedules")
    glazing_requirements = _extract_glazing_requirements(schedule_pages)
    (job_dir / "02_Schedules" / "glazing_thermal_requirements.json").write_text(
        json.dumps(glazing_requirements, indent=2) + "\n",
        encoding="utf-8",
    )
    structured_quote_rows = quote_rows_from_structured(structured_result)
    structured_items = [
        item for schedule in structured_result.schedules for item in schedule.items
    ]
    structured_uncertain_items = [
        item for item in structured_items if item.confidence < HUMAN_REVIEW_THRESHOLD
    ]

    glazing_rows = _dedupe_rows(
        [
            row
            for pdf_path in schedule_pages
            for row in _extract_glazing_rows(pdf_path)
        ],
        GLAZING_HEADERS[:-1],
    )
    windows = _dedupe_rows(
        structured_quote_rows["windows"]
        + [
            row
            for pdf_path in schedule_pages
            for row in _extract_windows(_extract_text(pdf_path), pdf_path.name)
        ]
        + [
            _window_from_glazing(row)
            for row in glazing_rows
            if "window" in row["description"].lower()
        ],
        WINDOW_HEADERS[:-1],
    )
    storefronts = _dedupe_rows(
        structured_quote_rows["storefronts"]
        + [
            row
            for pdf_path in schedule_pages
            for row in _extract_storefronts(_extract_text(pdf_path), pdf_path.name)
        ]
        + [
            _storefront_from_glazing(row)
            for row in glazing_rows
            if "storefront" in row["description"].lower()
        ],
        STOREFRONT_HEADERS[:-1],
    )
    structured_door_numbers = {
        row["door_no"] for row in structured_quote_rows["doors"] if row.get("door_no")
    }
    legacy_door_rows = [
        row
        for pdf_path in door_pages
        for row in _extract_doors(pdf_path)
        if row.get("door_no") not in structured_door_numbers
    ]
    doors = _dedupe_rows(
        structured_quote_rows["doors"]
        + legacy_door_rows
        + [
            _door_from_glazing(row)
            for row in glazing_rows
            if "door" in row["description"].lower()
        ],
        DOOR_HEADERS[:-1],
    )
    pressure_texts = [
        (pdf_path.name, _extract_text(pdf_path)) for pdf_path in pressure_pages
    ]
    zones = _extract_zone_tokens(pressure_texts)
    noas = _dedupe_rows(
        [
            row
            for pdf_path in schedule_pages
            for row in _extract_noas(_extract_text(pdf_path), pdf_path.name)
        ],
        ("noa", "expires", "scope", "status"),
    )
    missing_sources = [
        label
        for pages, label in (
            (door_pages, "door schedule"),
            (window_pages, "window schedule"),
            (
                [*storefront_pages, *structured_quote_rows["storefronts"]],
                "storefront or glazing schedule",
            ),
            (pressure_pages, "wind-pressure sheet"),
        )
        if not pages
    ]
    extra_qa_issues = [
        *_workbook_capacity_issues(
            template_path,
            structured_quote_rows["windows"],
            structured_quote_rows["storefronts"],
            structured_quote_rows["doors"],
        ),
        *_identified_source_extraction_issues(
            door_pages,
            window_pages,
            storefront_pages,
            pressure_pages,
            doors,
            windows,
            storefronts,
            zones,
        ),
        *_structured_extraction_issues(schedule_pages, structured_result),
        *_missing_workbook_attribute_issues(
            structured_quote_rows["windows"],
            structured_quote_rows["storefronts"],
            structured_quote_rows["doors"],
        ),
    ]
    if structured_quote_rows["storefronts"] and not storefront_pages:
        extra_qa_issues.append(
            {
                "severity": "MEDIUM",
                "issue_type": "Glazing schedule mapped as project storefront package",
                "mark": f"{len(structured_quote_rows['storefronts'])} glazing rows",
                "source_page": "02_Schedules/extracted_schedules.json",
                "description": "The project-specific glazing schedule was preserved together in drawing order on the Storefronts worksheet.",
                "recommended_action": "Review each glazing mark against the drawing before pricing; do not combine rows with different size, type, PSF, zone, or notes.",
            }
        )

    critical_issues, qa_issues = _write_reports(
        job_dir,
        project_name,
        address,
        windows,
        storefronts,
        doors,
        glazing_rows,
        zones,
        noas,
        missing_sources,
        extra_qa_issues,
    )
    draft_path = job_dir / "04_Quote_Workbook" / "filled_quote_workbook.xlsx"
    _create_workbook_draft(
        template_path,
        draft_path,
        structured_quote_rows["windows"],
        structured_quote_rows["storefronts"],
        structured_quote_rows["doors"],
        critical_issues,
        glazing_requirements,
    )
    _write_organized_plan_downloads(job_dir)
    _write_completed_checklist(job_dir, len(critical_issues))
    summary_path = _write_pipeline_summary(
        job_dir,
        windows,
        storefronts,
        doors,
        structured_quote_rows,
        glazing_rows,
        zones,
        noas,
        critical_issues,
        qa_issues,
        len(structured_items),
        len(structured_uncertain_items),
        project_name,
        address,
        notes,
        municipality_result,
    )
    _write_job_zip(job_dir)
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    if job_dir.parent.resolve() == JOBS_DIR.resolve():
        _append_jobs_index(job_dir, summary)
    print(f"[HurricaneOps] Created quote workbook draft: {draft_path}")
    print(f"[HurricaneOps] Created pipeline summary: {summary_path}")
    print(f"[HurricaneOps] CRITICAL issues found: {len(critical_issues)}")
    print(f"[HurricaneOps] Job folder: {job_dir}")
    return job_dir, critical_issues
