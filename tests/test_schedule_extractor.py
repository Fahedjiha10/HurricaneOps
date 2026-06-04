from __future__ import annotations

import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook


sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "05_SCRIPTS"))

from schedule_extractor import ScheduleExtractor
from bid_pipeline import _create_workbook_draft, _parse_glazing_requirements
from export_quote import quote_rows_from_structured
from schemas import DoorSchedule, ExtractionResult, GlazingSchedule


class ScheduleExtractorRowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.extractor = ScheduleExtractor(project_address="160 NW 44 St")

    def test_sample_glazing_rows(self) -> None:
        swing = self.extractor._glazing_item_from_cells(
            ["G01", "SWING DOOR", "1", "6'-0\"", "9'-0\"", "", "FL 26942", "MR. GLASS"],
            "LEVEL 1",
        )
        sliding = self.extractor._glazing_item_from_cells(
            ["G07", "SLIDING DOOR - 4 PANEL", "1", "15'-6\"", "10'-0\"", "", "FL 19092", "MR. GLASS"],
            "LEVEL 1",
        )
        self.assertIsNotNone(swing)
        self.assertIsNotNone(sliding)
        self.assertEqual(swing.width_inches, 72.0)
        self.assertEqual(sliding.width_inches, 186.0)
        self.assertGreaterEqual(swing.confidence, 0.85)
        self.assertGreaterEqual(sliding.confidence, 0.85)

    def test_sample_door_row(self) -> None:
        door = self.extractor._door_item_from_cells(
            ["101", "GARAGE", "1", "2'-8\"", "9'-0\"", "1", "", "WOOD", "SWING"],
            "LEVEL 1",
        )
        self.assertIsNotNone(door)
        self.assertEqual(door.door_number, "101")
        self.assertEqual(door.width_inches, 32.0)
        self.assertGreaterEqual(door.confidence, 0.85)

    def test_quote_adapter_keeps_glazing_package_together(self) -> None:
        garage = self.extractor._garage_door_items_from_text(
            "GARAGE DOOR SCHEDULE\nG1 GARAGE 1 9' - 0\" 9' - 0\" 1 METAL OVERHEAD METAL"
        )[0]
        glazing_door = self.extractor._glazing_item_from_cells(
            ["G01", "SWING DOOR", "1", "6'-0\"", "9'-0\"", "", "FL 26942", "MR. GLASS"],
            "LEVEL 1",
        )
        result = ExtractionResult(
            schedules=[
                GlazingSchedule("160 NW 44 St", 1, [glazing_door]),
                DoorSchedule("160 NW 44 St", 1, [garage]),
            ]
        )
        rows = quote_rows_from_structured(result)
        self.assertEqual([row["door_no"] for row in rows["doors"]], ["G1"])
        self.assertEqual([row["mark"] for row in rows["storefronts"]], ["G01"])
        self.assertEqual(rows["windows"], [])

    def test_quote_adapter_preserves_architectural_door_order(self) -> None:
        door = self.extractor._door_item_from_cells(
            ["101", "GARAGE", "1", "2'-8\"", "9'-0\"", "1", "", "WOOD", "SWING"],
            "LEVEL 1",
        )
        garage = self.extractor._garage_door_items_from_text(
            "GARAGE DOOR SCHEDULE\nG1 GARAGE 1 9' - 0\" 9' - 0\" 1 METAL OVERHEAD METAL BY MANUF. NOA REQUIRED, COORDINATE DOOR OPENING"
        )[0]
        rows = quote_rows_from_structured(
            ExtractionResult(schedules=[DoorSchedule("160 NW 44 St", 1, [door, garage])])
        )["doors"]
        self.assertEqual([row["door_no"] for row in rows], ["101", "G1"])
        self.assertEqual(rows[1]["door_material"], "METAL")
        self.assertEqual(rows[1]["panic_hardware"], "BY MANUF.")
        self.assertEqual(rows[1]["remarks"], "NOA REQUIRED, COORDINATE DOOR OPENING")

    def test_quote_adapter_copies_window_subset_without_splitting_glazing_package(self) -> None:
        window = self.extractor._glazing_item_from_cells(
            ["G02", "SLIDING WINDOW - 3 PANEL", "1", "9'-0\"", "5'-2\"", "", "FL 20359", "MR. GLASS"],
            "LEVEL 1",
        )
        rows = quote_rows_from_structured(
            ExtractionResult(schedules=[GlazingSchedule("160 NW 44 St", 1, [window])])
        )
        self.assertEqual([row["mark"] for row in rows["storefronts"]], ["G02"])
        self.assertEqual([row["mark"] for row in rows["windows"]], ["G02"])
        self.assertEqual(rows["windows"][0]["panels"], "3")
        self.assertEqual(rows["windows"][0]["glass_type"], "MR. GLASS")

    def test_quote_adapter_splits_door_frame_finish_from_combined_material(self) -> None:
        door = self.extractor._door_item_from_cells(
            [
                "101",
                "GARAGE",
                "1",
                "2'-8\"",
                "9'-0\"",
                "1",
                "",
                "WOOD",
                "SWING",
                "PAINTED WOOD, PAINTED WOOD",
            ],
            "LEVEL 1",
        )
        rows = quote_rows_from_structured(
            ExtractionResult(schedules=[DoorSchedule("160 NW 44 St", 1, [door])])
        )["doors"]
        self.assertEqual(rows[0]["door_material"], "PAINTED WOOD")
        self.assertEqual(rows[0]["frame_material"], "WOOD")
        self.assertEqual(rows[0]["frame_finish"], "PAINTED WOOD")

    def test_shifted_architectural_door_schedule_row_exports_key_fields(self) -> None:
        door = self.extractor._door_item_from_cells(
            [
                "",
                "001",
                "10'-0\"",
                "10'-0\"",
                "3\"",
                "D",
                "METAL",
                "PNT",
                "**",
                "",
                "METAL",
                "BAY ENTRANCE N",
                "OA No. 20-0417.04",
            ],
            None,
        )
        self.assertIsNotNone(door)
        self.assertEqual(door.door_number, "001")
        self.assertEqual(door.location, "BAY ENTRANCE")
        self.assertEqual(door.width_inches, 120.0)
        self.assertEqual(door.height_inches, 120.0)
        self.assertEqual(door.thickness, "3\"")
        self.assertEqual(door.jamb, "METAL")
        self.assertEqual(door.frame_finish, "PNT")
        self.assertEqual(door.noa, "NOA No. 20-0417.04")
        self.assertGreaterEqual(door.confidence, 0.85)

        rows = quote_rows_from_structured(
            ExtractionResult(schedules=[DoorSchedule("215 NW 63rd St", 1, [door])])
        )["doors"]
        self.assertEqual(rows[0]["door_no"], "001")
        self.assertEqual(rows[0]["location"], "BAY ENTRANCE")
        self.assertEqual(rows[0]["thickness"], "3\"")
        self.assertEqual(rows[0]["door_material"], "METAL")
        self.assertEqual(rows[0]["frame_material"], "METAL")
        self.assertEqual(rows[0]["frame_finish"], "PNT")
        self.assertEqual(rows[0]["noa"], "NOA No. 20-0417.04")
        self.assertEqual(rows[0]["remarks"], "")

    def test_multiline_architectural_door_rows_split_before_parsing(self) -> None:
        result = self.extractor._extract_table_rows(
            [
                [
                    [
                        "",
                        "021\n022",
                        "10'-0\"\n3'-0\"",
                        "10'-0\"\n8'-6\"",
                        "3\"\n1 3/4\"",
                        "D\nC",
                        "METAL\nALUM.",
                        "PNT\nPNT",
                        "**\n**",
                        "",
                        "METAL\nALUM.",
                        "BAY ENTRANCE N\nBAY ENTRANCE F",
                        "OA No. 20-0417.04\nLPA # FL15712-R3",
                    ]
                ]
            ],
            "",
            Path("A-8-01-DOOR-SCHEDULES.pdf"),
            1,
            "pdfplumber",
        )
        rows = quote_rows_from_structured(result)["doors"]
        self.assertEqual([row["door_no"] for row in rows], ["021", "022"])
        self.assertEqual(rows[0]["frame_material"], "METAL")
        self.assertEqual(rows[1]["frame_material"], "ALUM.")
        self.assertEqual(rows[1]["noa"], "FLPA # FL15712-R3")

    def test_glazing_thermal_requirements_are_parsed_without_guessing(self) -> None:
        requirements = _parse_glazing_requirements(
            "Glazing min. thermal standards: U-Factor = 1.08, SHGC = 0.45",
            "A-7.0",
        )
        self.assertEqual(
            requirements,
            {"u_factor": "1.08", "shgc": "0.45", "source": "A-7.0"},
        )

    def test_storefront_package_reconciles_explicit_thermal_requirements(self) -> None:
        storefront = {
            "mark": "G01",
            "quantity": "1",
            "width": "6'-0\"",
            "height": "9'-0\"",
            "type": "SWING DOOR",
            "material": "ALUMINUM",
            "glass_type": "MR. GLASS, SERIES MG-3000",
            "finish": "BRONZE",
            "noa": "FL 26942",
            "brand_product": "MR. GLASS, SERIES MG-3000",
            "level": "LEVEL 1",
            "remarks": "ENTRY, LATCH & DEADBOLT HARDWARE",
            "panels": "",
            "source_schedule": "GLAZING SCHEDULE",
            "extraction_status": "STRUCTURED VALIDATED - glazing schedule",
            "source": "extracted_schedules.json",
        }
        with TemporaryDirectory() as temp_dir:
            draft_path = Path(temp_dir) / "04_Quote_Workbook" / "draft.xlsx"
            draft_path.parent.mkdir()
            schedules_dir = Path(temp_dir) / "02_Schedules"
            schedules_dir.mkdir()
            _create_workbook_draft(
                Path(__file__).resolve().parents[1] / "00_INBOX" / "Template for Quotes .xlsx",
                draft_path,
                [],
                [storefront],
                [],
                [],
                {"u_factor": "1.08", "shgc": "0.45", "source": "A-7.0"},
            )
            workbook = load_workbook(draft_path, data_only=False)
            sheet = workbook["Storefronts"]
            self.assertEqual(sheet["A1"].value, "STOREFRONT GLAZING PACKAGE - UNIT COLUMNS")
            self.assertEqual(sheet["B4"].value, "G01")
            self.assertEqual(sheet["B10"].value, "ALUMINUM")
            self.assertEqual(sheet["B11"].value, "MR. GLASS, SERIES MG-3000")
            self.assertEqual(sheet["B12"].value, "BRONZE")
            self.assertEqual(sheet["B13"].value, "1.08")
            self.assertEqual(sheet["B14"].value, "0.45")
            self.assertEqual(sheet["A19"].value, "Count Check")
            self.assertEqual(sheet["A21"].value, "G01")
            self.assertEqual(sheet["B21"].value, 1)
            self.assertEqual(sheet["C21"].value, "=COUNTIF($B$4:$L$4,A21)")
            audit = (schedules_dir / "storefront_workbook_transfer_audit.json").read_text()
            self.assertIn('"passed": true', audit)

    def test_workbook_expands_and_reconciles_full_door_schedule(self) -> None:
        doors = []
        for number in range(101, 126):
            door_type = "SWING"
            if number == 103:
                door_type = "POCKET"
            elif number == 107:
                door_type = "BARN"
            doors.append(
                {
                    "door_no": str(number),
                    "location": f"ROOM {number}",
                    "type": door_type,
                    "width": "2' - 8\"",
                    "height": "7' - 0\"",
                    "thickness": "",
                    "door_material": "PAINTED WOOD",
                    "door_finish": "",
                    "frame_material": "WOOD",
                    "frame_finish": "PAINTED WOOD",
                    "fire_rating": "",
                    "noa": "",
                    "panic_hardware": "PRIVACY LOCK",
                    "remarks": "",
                    "level": "LEVEL 1",
                    "quantity": "1",
                    "panels": "1",
                    "fixed_panels": "",
                    "scope": "door schedule",
                    "extraction_status": "STRUCTURED VALIDATED - door schedule",
                    "raw_text": "",
                    "source": "extracted_schedules.json",
                }
            )
        with TemporaryDirectory() as temp_dir:
            draft_path = Path(temp_dir) / "04_Quote_Workbook" / "draft.xlsx"
            draft_path.parent.mkdir()
            (Path(temp_dir) / "02_Schedules").mkdir()
            _create_workbook_draft(
                Path(__file__).resolve().parents[1] / "00_INBOX" / "Template for Quotes .xlsx",
                draft_path,
                [],
                [],
                doors,
                [],
            )
            workbook = load_workbook(draft_path, data_only=False)
            sheet = workbook["Doors"]
            self.assertEqual(sheet["A1"].value, "DOOR SCHEDULE - UNIT COLUMNS")
            self.assertEqual(sheet["B4"].value, "101")
            self.assertEqual(sheet["B13"].value, "PAINTED WOOD")
            self.assertEqual(sheet["Z4"].value, "125")
            self.assertEqual(sheet["A20"].value, "Count Check")
            self.assertEqual(sheet["A21"].value, "Type")
            self.assertEqual(sheet["B21"].value, "Count")
            self.assertEqual(sheet["A22"].value, "SWING")
            self.assertEqual(sheet["B22"].value, "=COUNTIF($B$5:$Z$5,A22)")
            self.assertEqual(sheet["A23"].value, "POCKET")
            self.assertEqual(sheet["B23"].value, "=COUNTIF($B$5:$Z$5,A23)")
            self.assertEqual(sheet["A24"].value, "BARN")
            self.assertEqual(sheet["B24"].value, "=COUNTIF($B$5:$Z$5,A24)")
            self.assertEqual(sheet["A25"].value, "Total : ")
            self.assertEqual(sheet["B25"].value, "=SUM(B22:B24)")
            self.assertIn("A1:Z1", [str(item) for item in sheet.merged_cells.ranges])

    def test_workbook_expands_and_reconciles_window_schedule(self) -> None:
        windows = []
        for number in range(1, 31):
            windows.append(
                {
                    "mark": f"G{number:02d}",
                    "quantity": "1",
                    "width": "3'-0\"",
                    "height": "5'-0\"",
                    "type": "FIXED WINDOW",
                    "material": "",
                    "glass_type": "MR. GLASS, SERIES MG-350",
                    "finish": "WHITE",
                    "noa": "FL 41889",
                    "brand_product": "MR. GLASS, SERIES MG-350",
                    "level": "LEVEL 1",
                    "remarks": "",
                    "panels": "",
                    "source_schedule": "GLAZING SCHEDULE",
                    "extraction_status": "STRUCTURED VALIDATED - glazing schedule",
                    "source": "extracted_schedules.json",
                }
            )
        with TemporaryDirectory() as temp_dir:
            draft_path = Path(temp_dir) / "04_Quote_Workbook" / "draft.xlsx"
            draft_path.parent.mkdir()
            (Path(temp_dir) / "02_Schedules").mkdir()
            _create_workbook_draft(
                Path(__file__).resolve().parents[1] / "00_INBOX" / "Template for Quotes .xlsx",
                draft_path,
                windows,
                [],
                [],
                [],
            )
            workbook = load_workbook(draft_path, data_only=False)
            sheet = workbook["Windows"]
            self.assertEqual(sheet["B6"].value, "G01")
            self.assertEqual(sheet["AE6"].value, "G30")
            self.assertEqual(sheet["B5"].value, "FIXED WINDOW")
            self.assertEqual(sheet["B7"].value, "GLAZING SCHEDULE")
            self.assertEqual(sheet["B11"].value, "MR. GLASS, SERIES MG-350")
            self.assertEqual(sheet["B12"].value, "WHITE")
            self.assertEqual(sheet["A20"].value, "Count Check")
            self.assertEqual(sheet["A22"].value, "G01")
            self.assertEqual(sheet["B22"].value, 1)
            self.assertEqual(sheet["C22"].value, "=COUNTIF($B$6:$AE$6,A22)")
            self.assertIn("A1:AE1", [str(item) for item in sheet.merged_cells.ranges])
