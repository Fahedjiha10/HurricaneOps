from __future__ import annotations

import csv
import hashlib
import io
import json
import subprocess
import sys
import tempfile
import time
from contextlib import redirect_stdout
from datetime import datetime
from pathlib import Path

import fitz
from openpyxl import load_workbook
import streamlit as st


PROJECT_ROOT = Path(__file__).resolve().parent
SCRIPTS_DIR = PROJECT_ROOT / "05_SCRIPTS"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from bid_pipeline import backfill_organized_plan_downloads, _write_job_zip, run_bid_pipeline
from config import (
    ICLOUD_AUTO_JOBS_DIR,
    ICLOUD_AUTO_JOBS_INDEX_PATH,
    ICLOUD_AUTOMATION_DIR,
    ICLOUD_INCOMING_DIR,
    ICLOUD_PROCESSED_DIR,
    ICLOUD_SYNC_STATE_PATH,
    INBOX_DIR,
    JOBS_DIR,
    JOBS_INDEX_PATH,
    ORGANIZED_PLAN_COMPARTMENTS,
    ORGANIZED_PLAN_SCHEMA_VERSION,
)
from job_intake import IntakeError, inspect_inbox
from municipality_lookup import lookup_municipality


REVIEW_STATUSES = ("Open", "Resolved", "Accepted Risk", "Denied / Needs RFI")

WORKBOOK_FIELD_ROWS = {
    "Windows": {
        "id_row": 6,
        "fields": {
            "Material": 10,
            "Glass Type": 11,
            "Frame Finish": 12,
            "U Factor": 14,
            "SHGC": 15,
            "PSF / Zone": 16,
            "Notes": 17,
        },
    },
    "Storefronts": {
        "id_row": 4,
        "fields": {
            "Material": 10,
            "Glass Type": 11,
            "Finish": 12,
            "U Factor": 13,
            "SHGC": 14,
            "PSF / Zone": 15,
            "Notes": 16,
        },
    },
    "Doors": {
        "id_row": 4,
        "fields": {
            "Door Material": 11,
            "Frame Material": 12,
            "Frame Finish Color": 13,
            "U Factor": 14,
            "SHGC": 15,
            "PSF / Zone": 16,
            "Notes": 17,
        },
    },
}


def load_summaries() -> list[dict[str, object]]:
    summaries: list[dict[str, object]] = []
    summary_paths = {
        summary_path
        for jobs_dir in (JOBS_DIR, ICLOUD_AUTO_JOBS_DIR)
        for summary_path in jobs_dir.glob("*/pipeline_summary.json")
    }
    for summary_path in sorted(summary_paths, reverse=True):
        summary = json.loads(summary_path.read_text(encoding="utf-8"))
        summary.setdefault("job_id", summary.get("job_name", summary_path.parent.name))
        summary.setdefault("project_name", "Legacy job")
        summary.setdefault("address", "Not recorded")
        summary.setdefault("municipality_location", "")
        summary.setdefault("municipality_lookup", {})
        summary.setdefault("qa_issues", [])
        summary.setdefault(
            "qa_issue_counts",
            {"LOW": 0, "MEDIUM": 0, "HIGH": 0, "CRITICAL": summary.get("critical_count", 0)},
        )
        summaries.append(summary)
    return sorted(summaries, key=lambda summary: str(summary["generated_at"]), reverse=True)


def load_job_history() -> list[dict[str, str]]:
    rows: dict[str, dict[str, str]] = {}
    for index_path in (JOBS_INDEX_PATH, ICLOUD_AUTO_JOBS_INDEX_PATH):
        if not index_path.exists():
            continue
        with index_path.open(newline="", encoding="utf-8") as handle:
            for row in csv.DictReader(handle):
                rows[row["job_id"]] = row
    return sorted(rows.values(), key=lambda row: row.get("created_at", ""), reverse=True)


def render_municipality_lookup(result: dict[str, object]) -> None:
    status = str(result.get("status") or "")
    display_name = str(result.get("display_name") or "")
    warning = str(result.get("warning") or "")
    if status == "FOUND" and display_name:
        st.success(f"Municipality / jurisdiction: {display_name}")
    elif status == "UNINCORPORATED_OR_UNKNOWN":
        st.warning(
            f"Municipality / jurisdiction needs review: {display_name or 'Not confirmed'}"
        )
    elif status in {"NO_MATCH", "LOOKUP_FAILED"}:
        st.error(warning or "Municipality lookup failed.")
    elif status == "NOT_PROVIDED":
        st.info(warning or "Enter an address to look up the municipality.")
    elif display_name:
        st.info(f"Municipality / jurisdiction: {display_name}")
    if result.get("matched_address"):
        st.caption(f"Matched address: {result['matched_address']}")
    if result.get("municipality_type"):
        st.caption(f"Lookup type: {result['municipality_type']} | Source: {result.get('source', '')}")
    if warning and status not in {"NO_MATCH", "LOOKUP_FAILED", "NOT_PROVIDED"}:
        st.caption(warning)


def load_icloud_sync_state() -> dict[str, object]:
    if not ICLOUD_SYNC_STATE_PATH.exists():
        return {"packages": {}}
    try:
        return json.loads(ICLOUD_SYNC_STATE_PATH.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {"packages": {}}


def incoming_icloud_packages() -> list[Path]:
    if not ICLOUD_INCOMING_DIR.exists():
        return []
    return sorted(path for path in ICLOUD_INCOMING_DIR.iterdir() if path.is_dir())


def run_icloud_queue_check(stable_seconds: int) -> tuple[bool, str]:
    automation_python = ICLOUD_AUTOMATION_DIR / ".venv" / "bin" / "python"
    automation_worker = ICLOUD_AUTOMATION_DIR / "05_SCRIPTS" / "icloud_auto_ingest.py"
    worker_python = automation_python if automation_python.exists() else Path(sys.executable)
    worker_script = (
        automation_worker
        if automation_worker.exists()
        else SCRIPTS_DIR / "icloud_auto_ingest.py"
    )
    try:
        result = subprocess.run(
            [
                str(worker_python),
                str(worker_script),
                "--once",
                "--stable-seconds",
                str(stable_seconds),
            ],
            capture_output=True,
            text=True,
            timeout=600,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        return False, f"iCloud queue check failed: {exc}"
    output = (result.stdout + result.stderr).strip()
    return (
        result.returncode == 0,
        output or "iCloud queue check completed with no terminal output.",
    )


def issue_id(issue: dict[str, str]) -> str:
    identity = "|".join(
        issue.get(field, "")
        for field in ("severity", "issue_type", "mark", "source_page", "description")
    )
    return hashlib.sha1(identity.encode("utf-8")).hexdigest()[:12]


def review_state_path(job_dir: Path) -> Path:
    return job_dir / "03_Takeoff_QA" / "estimator_review_state.json"


def load_review_state(job_dir: Path, issues: list[dict[str, str]]) -> dict[str, object]:
    path = review_state_path(job_dir)
    saved = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    saved_issues = saved.get("issues", {})
    current_issues: dict[str, dict[str, str]] = {}
    for issue in issues:
        current_id = issue_id(issue)
        prior = saved_issues.get(current_id, {})
        current_issues[current_id] = {
            "status": prior.get("status", "Open"),
            "estimator_note": prior.get("estimator_note", ""),
            "resolved_by": prior.get("resolved_by", ""),
            "resolution_date": prior.get("resolution_date", ""),
        }
    return {
        "issues": current_issues,
        "manual_entries": saved.get("manual_entries", []),
        "approved_for_send": bool(saved.get("approved_for_send", False)),
        "approved_by": saved.get("approved_by", ""),
        "approved_at": saved.get("approved_at", ""),
    }


def save_review_state(job_dir: Path, state: dict[str, object]) -> None:
    path = review_state_path(job_dir)
    path.write_text(json.dumps(state, indent=2) + "\n", encoding="utf-8")


def _join_note_parts(*parts: str) -> str:
    return "; ".join(part.strip() for part in parts if part and part.strip())


def approval_decision_path(job_dir: Path) -> Path:
    return job_dir / "approval_decision.json"


def load_approval_decision(job_dir: Path) -> dict[str, object]:
    path = approval_decision_path(job_dir)
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}


def save_risk_acceptance(
    job_dir: Path, critical_issue_count: int, estimator_notes: str
) -> None:
    decision = {
        "approved_with_risk": True,
        "approved_at": datetime.now().isoformat(timespec="seconds"),
        "critical_issue_count": critical_issue_count,
        "estimator_notes": estimator_notes.strip(),
        "qa_report_path": str((job_dir / "03_Takeoff_QA" / "qa_report.xlsx").resolve()),
        "psf_zone_report_path": str(
            (job_dir / "05_NOA_PSF_Check" / "psf_zone_report.xlsx").resolve()
        ),
    }
    approval_decision_path(job_dir).write_text(
        json.dumps(decision, indent=2) + "\n", encoding="utf-8"
    )


def unresolved_critical_count(issues: list[dict[str, str]], state: dict[str, object]) -> int:
    return sum(
        issue.get("severity") == "CRITICAL"
        and state["issues"][issue_id(issue)]["status"] == "Open"
        for issue in issues
    )


def ensure_pdf_preview(pdf_path: Path) -> Path:
    preview_dir = pdf_path.parent / "previews"
    preview_dir.mkdir(exist_ok=True)
    preview_path = preview_dir / f"{pdf_path.stem}.png"
    valid_preview = (
        preview_path.exists()
        and preview_path.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"
        and preview_path.stat().st_mtime >= pdf_path.stat().st_mtime
    )
    if not valid_preview:
        with tempfile.NamedTemporaryFile(
            dir=preview_dir, suffix=".png", delete=False
        ) as temp_file:
            temp_path = Path(temp_file.name)
        with fitz.open(pdf_path) as document:
            pixmap = document[0].get_pixmap(matrix=fitz.Matrix(0.55, 0.55), alpha=False)
            pixmap.save(temp_path)
        temp_path.replace(preview_path)
    return preview_path


def resolve_review_documents(job_dir: Path, source_page: str) -> list[Path]:
    documents: list[Path] = []
    organized_dir = job_dir / "07_Organized_Plan_Set"

    def add_compartment_documents(directory_name: str) -> None:
        compartment_dir = organized_dir / directory_name
        if compartment_dir.exists():
            documents.extend(sorted(path for path in compartment_dir.glob("*.pdf") if path.is_file()))

    tokens = [
        token.strip()
        for token in source_page.replace(";", ",").split(",")
        if token.strip()
    ]
    for token in tokens:
        candidate = job_dir / token
        if candidate.exists():
            documents.append(candidate)
            continue
        token_path = Path(token)
        if token_path.suffix:
            matches = sorted(
                path
                for path in job_dir.rglob(token_path.name)
                if path.is_file() and "previews" not in path.parts
            )
            documents.extend(matches)
            continue
        lowered = token.lower()
        if "worksheet" in lowered:
            workbook = job_dir / "04_Quote_Workbook" / "filled_quote_workbook.xlsx"
            if workbook.exists():
                documents.append(workbook)
        elif "indexed" in lowered or "sheet" in lowered:
            sheet_index = job_dir / "07_Organized_Plan_Set" / "sheet_index.csv"
            if sheet_index.exists():
                documents.append(sheet_index)
    lowered_source = source_page.lower()
    if "indexed schedule sheets" in lowered_source:
        add_compartment_documents("02_Door_Schedules")
        add_compartment_documents("03_Window_Schedules")
        add_compartment_documents("04_Storefront_Schedules")
    if "door" in lowered_source and "schedule" in lowered_source:
        add_compartment_documents("02_Door_Schedules")
    if "window" in lowered_source and "schedule" in lowered_source:
        add_compartment_documents("03_Window_Schedules")
    if "storefront" in lowered_source and "schedule" in lowered_source:
        add_compartment_documents("04_Storefront_Schedules")
    if "wind" in lowered_source and ("pressure" in lowered_source or "psf" in lowered_source):
        add_compartment_documents("06_Wind_Pressure_Elevations")
        psf_report = job_dir / "05_NOA_PSF_Check" / "psf_zone_report.xlsx"
        if psf_report.exists():
            documents.append(psf_report)
    if not documents:
        fallback = job_dir / "03_Takeoff_QA" / "qa_report.xlsx"
        if fallback.exists():
            documents.append(fallback)
    unique_documents: list[Path] = []
    seen: set[Path] = set()
    for document in documents:
        resolved = document.resolve()
        if resolved in seen:
            continue
        seen.add(resolved)
        unique_documents.append(document)
    return unique_documents


def render_review_documents(job_dir: Path, issue: dict[str, str], key_prefix: str) -> None:
    documents = resolve_review_documents(job_dir, issue.get("source_page", ""))
    st.write("**Linked review document(s)**")
    if not documents:
        st.info("No source document could be linked automatically. Use the QA report and organized plan set.")
        return
    for index, document in enumerate(documents):
        document_key = hashlib.sha1(str(document.resolve()).encode("utf-8")).hexdigest()[:10]
        st.caption(str(document.relative_to(job_dir) if document.is_relative_to(job_dir) else document))
        if document.suffix.lower() == ".pdf":
            try:
                st.image(ensure_pdf_preview(document), caption=document.name, width="stretch")
            except (OSError, ValueError, RuntimeError) as exc:
                st.warning(f"Could not preview {document.name}: {exc}")
            mime = "application/pdf"
        elif document.suffix.lower() == ".xlsx":
            mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        elif document.suffix.lower() == ".csv":
            mime = "text/csv"
        else:
            mime = "application/octet-stream"
        st.download_button(
            f"Download/open {document.name}",
            data=document.read_bytes(),
            file_name=document.name,
            mime=mime,
            key=f"{key_prefix}-doc-{index}-{document_key}",
            width="stretch",
        )


def infer_workbook_tab(issue: dict[str, str]) -> str:
    text = " ".join(
        issue.get(field, "")
        for field in ("issue_type", "mark", "source_page", "description", "recommended_action")
    ).lower()
    if "storefront" in text or "glazing" in text:
        return "Storefronts"
    if "window" in text:
        return "Windows"
    if "door" in text:
        return "Doors"
    return "Windows"


def clean_issue_mark_for_workbook(mark: str) -> str:
    stripped = mark.strip()
    if not stripped or stripped.lower() in {"all openings", "psf / zone values"}:
        return ""
    if stripped.lower().endswith(" rows"):
        return ""
    return stripped


def apply_manual_workbook_entry(
    job_dir: Path,
    tab_name: str,
    item_id: str,
    field_name: str,
    value: str,
    overwrite: bool,
) -> tuple[Path, str]:
    if tab_name not in WORKBOOK_FIELD_ROWS:
        raise ValueError(f"Unknown workbook tab: {tab_name}")
    spec = WORKBOOK_FIELD_ROWS[tab_name]
    field_rows = spec["fields"]
    if field_name not in field_rows:
        raise ValueError(f"Unknown workbook field for {tab_name}: {field_name}")
    workbook_path = job_dir / "04_Quote_Workbook" / "filled_quote_workbook.xlsx"
    if not workbook_path.exists():
        raise FileNotFoundError(f"Missing filled quote workbook: {workbook_path}")
    workbook = load_workbook(workbook_path, data_only=False)
    sheet = workbook[tab_name]
    normalized_item_id = item_id.strip().upper()
    target_column = None
    for column_index in range(2, sheet.max_column + 1):
        cell_value = str(sheet.cell(spec["id_row"], column_index).value or "").strip().upper()
        if cell_value == normalized_item_id:
            target_column = column_index
            break
    if target_column is None:
        raise ValueError(f"Could not find `{item_id}` on the {tab_name} worksheet.")
    target_cell = sheet.cell(field_rows[field_name], target_column)
    if field_name == "Notes" and target_cell.value:
        target_cell.value = f"{target_cell.value}; Manual review: {value.strip()}"
    else:
        existing = str(target_cell.value or "").strip()
        if existing and existing != value.strip() and not overwrite:
            raise ValueError(
                f"{tab_name}!{target_cell.coordinate} already contains `{existing}`. "
                "Enable overwrite to replace it."
            )
        target_cell.value = value.strip()
    workbook.save(workbook_path)
    return workbook_path, f"{tab_name}!{target_cell.coordinate}"


def render_download(
    job_dir: Path, label: str, relative_path: str, disabled: bool = False
) -> None:
    path = job_dir / relative_path
    if not path.exists():
        st.warning(f"Missing output: {relative_path}")
        return
    st.download_button(
        label=label,
        data=path.read_bytes(),
        file_name=path.name,
        mime="application/octet-stream",
        key=f"download-{job_dir.name}-{path.name}",
        width="stretch",
        disabled=disabled,
    )


def stage_uploads(pdf_files, excel_file) -> tempfile.TemporaryDirectory[str]:
    temp_dir = tempfile.TemporaryDirectory(prefix="hurricaneops_upload_")
    inbox_dir = Path(temp_dir.name) / "00_INBOX"
    plan_set_dir = inbox_dir / "Uploaded_Plan_Set"
    plan_set_dir.mkdir(parents=True)
    for pdf_file in pdf_files:
        (plan_set_dir / Path(pdf_file.name).name).write_bytes(pdf_file.getvalue())
    (inbox_dir / Path(excel_file.name).name).write_bytes(excel_file.getvalue())
    return temp_dir


st.set_page_config(page_title="HurricaneOps Estimator Dashboard", page_icon="H", layout="wide")
st.title("HurricaneOps Estimator Review Dashboard")
st.caption("Local-only bid intake, takeoff review, QA gating, and estimator download package.")

@st.fragment(run_every=120)
def render_icloud_automation() -> None:
    now = time.monotonic()
    last_auto_check = float(st.session_state.get("icloud_last_auto_check", 0))
    if now - last_auto_check >= 115:
        success, output = run_icloud_queue_check(stable_seconds=60)
        st.session_state["icloud_sync_success"] = success
        st.session_state["icloud_sync_output"] = output
        st.session_state["icloud_last_auto_check"] = now

    st.subheader("iCloud Automation")
    with st.container(border=True):
        incoming_packages = incoming_icloud_packages()
        sync_state = load_icloud_sync_state()
        package_records = list(sync_state.get("packages", {}).values())
        queue_left, queue_right = st.columns(2)
        with queue_left:
            st.write(f"**Incoming queue:** `{ICLOUD_INCOMING_DIR}`")
            st.write(f"**Processed results:** `{ICLOUD_PROCESSED_DIR}`")
            st.metric("Waiting project folders", len(incoming_packages))
        with queue_right:
            st.write(
                "Drop each new bid into its own project folder inside the incoming queue. "
                "Include one combined plan-set PDF or the separate plan PDFs; the standard "
                "iCloud quote template is used automatically unless the project folder includes "
                "one Excel workbook."
            )
            st.caption(
                "While this local dashboard is open, it checks every two minutes. Files must "
                "remain unchanged for one minute before an automatic run begins."
            )
            check_queue = st.button("Check iCloud queue now", width="stretch")

        if check_queue:
            success, output = run_icloud_queue_check(stable_seconds=0)
            st.session_state["icloud_sync_success"] = success
            st.session_state["icloud_sync_output"] = output
            st.session_state["icloud_last_auto_check"] = time.monotonic()
            st.rerun(scope="fragment")

        if "icloud_sync_output" in st.session_state:
            if st.session_state.get("icloud_sync_success"):
                st.success("iCloud queue check completed.")
            else:
                st.error("iCloud queue check reported an error.")
            with st.expander("iCloud queue terminal output"):
                st.code(st.session_state["icloud_sync_output"], language="text")

        if package_records:
            st.caption("Recent iCloud package activity")
            st.dataframe(
                [
                    {
                        "Project folder": Path(str(record.get("source", ""))).name,
                        "Status": record.get("status", ""),
                        "Updated": record.get("updated_at", ""),
                        "Critical issues": record.get("critical_count", ""),
                        "Generated job": Path(str(record.get("job_dir", ""))).name,
                        "Error": record.get("error", ""),
                    }
                    for record in sorted(
                        package_records,
                        key=lambda record: str(record.get("updated_at", "")),
                        reverse=True,
                    )
                ],
                hide_index=True,
                width="stretch",
            )


render_icloud_automation()

st.subheader("New Job Intake")
with st.container(border=True):
    project_left, project_right = st.columns(2)
    project_name = project_left.text_input("Project name", placeholder="Example: 268 Co-Living")
    project_address = project_right.text_input(
        "Project address", placeholder="Example: 268 NE 80th Ter, Miami FL"
    )
    lookup_key = "new_job_municipality_lookup"
    lookup_address = st.session_state.get(f"{lookup_key}_address", "")
    if project_right.button(
        "Find municipality",
        disabled=not project_address.strip(),
        help="Uses the entered address to look up the incorporated municipality or jurisdiction.",
        width="stretch",
    ):
        with st.spinner("Looking up municipality..."):
            st.session_state[lookup_key] = lookup_municipality(project_address).to_dict()
            st.session_state[f"{lookup_key}_address"] = project_address.strip()
            lookup_address = project_address.strip()
    municipality_preview = st.session_state.get(lookup_key)
    if (
        isinstance(municipality_preview, dict)
        and lookup_address == project_address.strip()
    ):
        render_municipality_lookup(municipality_preview)
    elif project_address.strip():
        st.caption("Municipality will also be looked up automatically when the pipeline runs.")
    notes = st.text_area("Optional notes", placeholder="Estimator notes, bid date, scope reminders...")
    use_existing_inbox = st.checkbox(
        "Use the files already staged in 00_INBOX",
        value=True,
        help="Turn this off to stage a new local upload set for this run.",
    )
    upload_left, upload_right = st.columns(2)
    pdf_files = upload_left.file_uploader(
        "Upload PDF plan set",
        type=["pdf"],
        accept_multiple_files=True,
        disabled=use_existing_inbox,
        help="Select one combined plan-set PDF or multiple PDF sheets. Files stay local to this Mac.",
    )
    excel_file = upload_right.file_uploader(
        "Upload Excel quote template",
        type=["xlsx", "xlsm", "xltx", "xltm", "xls"],
        disabled=use_existing_inbox,
        help="Select one quote workbook template. Files stay local to this Mac.",
    )

    staged_status = inspect_inbox()
    intake_ready = bool(project_name.strip())
    if use_existing_inbox:
        intake_ready = intake_ready and bool(staged_status["ready"])
        st.caption(
            f"Staged inbox: {staged_status['pdf_count']} source PDF files and "
            f"`{getattr(staged_status['excel_template'], 'name', 'no workbook')}`"
        )
    else:
        intake_ready = intake_ready and bool(pdf_files) and excel_file is not None
        st.caption(f"Uploaded plan-set PDF files: {len(pdf_files or [])}")

    run_clicked = st.button(
        "Run bid pipeline",
        type="primary",
        disabled=not intake_ready,
        width="stretch",
    )

if run_clicked:
    output = io.StringIO()
    upload_temp_dir = None
    try:
        pipeline_inbox = INBOX_DIR
        if not use_existing_inbox:
            upload_temp_dir = stage_uploads(pdf_files, excel_file)
            pipeline_inbox = Path(upload_temp_dir.name) / "00_INBOX"
        with st.spinner("Running HurricaneOps..."), redirect_stdout(output):
            job_dir, issues = run_bid_pipeline(
                inbox_dir=pipeline_inbox,
                project_name=project_name.strip(),
                address=project_address.strip(),
                notes=notes.strip(),
            )
    except (IntakeError, OSError) as exc:
        st.error(f"Pipeline failed: {exc}")
    else:
        st.success(f"Pipeline completed: {job_dir.name}")
        if issues:
            st.warning(f"{len(issues)} CRITICAL review items require attention.")
        else:
            st.success("No CRITICAL issues found.")
    finally:
        if upload_temp_dir is not None:
            upload_temp_dir.cleanup()
    with st.expander("Terminal output"):
        st.code(output.getvalue() or "No terminal output captured.", language="text")

summaries = load_summaries()
st.divider()
st.subheader("Job Status Dashboard")
if not summaries:
    st.info("No completed estimator jobs found yet.")
else:
    selected_summary = summaries[0]
    if len(summaries) > 1:
        selected_job = st.selectbox(
            "Dashboard job",
            options=[str(summary["job_id"]) for summary in summaries],
            index=0,
        )
        selected_summary = next(
            summary for summary in summaries if summary["job_id"] == selected_job
        )

    job_dir = Path(str(selected_summary["job_dir"]))
    counts = selected_summary["counts"]
    qa_issues = selected_summary.get("qa_issues", [])
    review_state = load_review_state(job_dir, qa_issues)
    approval_decision = load_approval_decision(job_dir)
    risk_accepted = bool(approval_decision.get("approved_with_risk", False))
    open_critical_count = unresolved_critical_count(qa_issues, review_state)
    issue_counts = selected_summary.get(
        "qa_issue_counts", {"LOW": 0, "MEDIUM": 0, "HIGH": 0, "CRITICAL": selected_summary["critical_count"]}
    )
    total_items = counts.get(
        "structured_schedule_rows",
        counts["windows"] + counts["storefronts"] + counts["common_area_doors"],
    )

    st.write(f"**Job name:** `{selected_summary['job_name']}`")
    st.write(f"**Job folder:** `{job_dir}`")
    st.write(f"**Created:** `{selected_summary['generated_at']}`")
    municipality_lookup = selected_summary.get("municipality_lookup", {})
    municipality_location = str(selected_summary.get("municipality_location") or "")
    if isinstance(municipality_lookup, dict) and municipality_lookup:
        render_municipality_lookup(municipality_lookup)
    elif municipality_location:
        st.write(f"**Municipality / jurisdiction:** `{municipality_location}`")
    else:
        st.info("Municipality / jurisdiction not recorded for this job.")

    item_metrics = st.columns(5)
    item_metrics[0].metric("Windows", counts.get("quote_window_rows", counts["windows"]))
    item_metrics[1].metric(
        "Storefront Glazing Package",
        counts.get("quote_glazing_storefront_rows", counts["storefronts"]),
    )
    item_metrics[2].metric("Doors", counts.get("quote_door_rows", counts["common_area_doors"]))
    item_metrics[3].metric("Glazing Schedule Rows", counts.get("glazing_schedule_rows", 0))
    item_metrics[4].metric("Total Extracted Items", total_items)

    qa_metrics = st.columns(4)
    qa_metrics[0].metric("LOW", issue_counts["LOW"])
    qa_metrics[1].metric("MEDIUM", issue_counts["MEDIUM"])
    qa_metrics[2].metric("HIGH", issue_counts["HIGH"])
    qa_metrics[3].metric("Open CRITICAL", open_critical_count, delta=f"{issue_counts['CRITICAL']} total")

    organized_source_counts = counts.get("organized_source_pdfs", {})
    organized_schema_version = int(selected_summary.get("organized_plan_schema_version", 0))
    organized_downloads_current = (
        organized_source_counts
        and organized_schema_version >= ORGANIZED_PLAN_SCHEMA_VERSION
    )
    if organized_downloads_current:
        st.subheader("Organized Plan Compartments")
        st.write(f"**Review folder:** `{job_dir / '07_Organized_Plan_Set'}`")
        st.write(
            f"**Indexed plan sheets:** `{counts.get('indexed_plan_sheets', 'Not recorded')}`"
        )
        organized_columns = st.columns(len(ORGANIZED_PLAN_COMPARTMENTS))
        for column, (key, _folder_name, label) in zip(
            organized_columns, ORGANIZED_PLAN_COMPARTMENTS
        ):
            column.metric(label, organized_source_counts.get(key, 0))
        st.caption(
            "The original plan-set hierarchy remains preserved in 00_Source_PDFs. "
            "The organized copies, source_file_manifest.csv, and sheet_index.csv are for estimator review."
        )
    else:
        with st.container(border=True):
            if organized_source_counts:
                st.write("**Organized plan downloads need the new elevation split.**")
                action_label = "Update organized downloads for this job"
            else:
                st.write("**Organized plan downloads are not available for this earlier job yet.**")
                action_label = "Create organized downloads for this job"
            st.caption(
                "Create separate architectural-elevation and wind-pressure-elevation ZIPs, "
                "along with floor-plan, schedule, and complete organized-plan downloads "
                "from the preserved source PDFs."
            )
            if st.button(
                action_label,
                key=f"{job_dir.name}-backfill-organized-downloads",
                width="stretch",
            ):
                try:
                    with st.spinner("Organizing preserved source PDFs..."):
                        backfill_organized_plan_downloads(job_dir)
                except (IntakeError, OSError, ValueError) as exc:
                    st.error(f"Could not create organized downloads: {exc}")
                else:
                    st.success("Organized downloads updated for this job.")
                    st.rerun()

    if open_critical_count:
        st.error("DO NOT SEND THIS QUOTE YET. Critical review items must be fixed first.")
    else:
        st.success("All CRITICAL issues have a resolution. Complete estimator approval before sending.")

    st.subheader("Critical Issues Panel")
    critical_rows = [
        issue
        for issue in qa_issues
        if issue["severity"] == "CRITICAL"
    ]
    if critical_rows:
        st.dataframe(
            [
                {
                    "Severity": row.get("severity", ""),
                    "Issue type": row.get("issue_type", ""),
                    "Mark": row.get("mark", ""),
                    "Sheet/source page": row.get("source_page", ""),
                    "Description": row.get("description", ""),
                    "Recommended action": row.get("recommended_action", ""),
                    "Resolution status": review_state["issues"][issue_id(row)]["status"],
                }
                for row in critical_rows
            ],
            hide_index=True,
            width="stretch",
        )
        st.subheader("Critical Issue Action Center")
        st.caption(
            "Each issue below links to the best matching review document. "
            "You can accept, deny/RFI, or manually enter a verified value into the filled quote workbook."
        )
        for index, issue in enumerate(critical_rows, start=1):
            current_issue_id = issue_id(issue)
            issue_state = review_state["issues"][current_issue_id]
            key_prefix = f"{job_dir.name}-{current_issue_id}"
            with st.expander(
                f"{index}. {issue.get('issue_type', 'Issue')} - {issue.get('mark', '')} "
                f"({issue_state['status']})",
                expanded=issue_state["status"] == "Open",
            ):
                st.write(f"**Problem:** {issue.get('description', '')}")
                st.write(f"**Recommended action:** {issue.get('recommended_action', '')}")
                render_review_documents(job_dir, issue, key_prefix)

                decision_columns = st.columns(3)
                with decision_columns[0]:
                    decision_status = st.selectbox(
                        "Review decision",
                        REVIEW_STATUSES,
                        index=REVIEW_STATUSES.index(issue_state["status"])
                        if issue_state["status"] in REVIEW_STATUSES
                        else 0,
                        key=f"{key_prefix}-decision",
                    )
                with decision_columns[1]:
                    resolved_by = st.text_input(
                        "Estimator",
                        value=issue_state.get("resolved_by", ""),
                        key=f"{key_prefix}-resolved-by",
                    )
                with decision_columns[2]:
                    resolution_date = st.text_input(
                        "Date",
                        value=issue_state.get(
                            "resolution_date",
                            datetime.now().date().isoformat()
                            if issue_state["status"] != "Open"
                            else "",
                        ),
                        key=f"{key_prefix}-resolution-date",
                    )
                estimator_note = st.text_area(
                    "Estimator note / reason",
                    value=issue_state.get("estimator_note", ""),
                    key=f"{key_prefix}-note",
                )
                if st.button(
                    "Save this issue decision",
                    key=f"{key_prefix}-save-decision",
                ):
                    review_state["issues"][current_issue_id] = {
                        "status": decision_status,
                        "estimator_note": estimator_note.strip(),
                        "resolved_by": resolved_by.strip(),
                        "resolution_date": resolution_date.strip(),
                    }
                    review_state["approved_for_send"] = False
                    review_state["approved_by"] = ""
                    review_state["approved_at"] = ""
                    save_review_state(job_dir, review_state)
                    _write_job_zip(job_dir)
                    st.success("Issue decision saved. Estimator approval was reset.")
                    st.rerun()

                st.write("**Manual workbook entry**")
                st.caption(
                    "Use this only after you verified the value in the linked document. "
                    "The entry is written into `filled_quote_workbook.xlsx`."
                )
                inferred_tab = infer_workbook_tab(issue)
                tab_name = st.selectbox(
                    "Workbook tab",
                    tuple(WORKBOOK_FIELD_ROWS),
                    index=tuple(WORKBOOK_FIELD_ROWS).index(inferred_tab),
                    key=f"{key_prefix}-manual-tab",
                )
                manual_columns = st.columns(3)
                with manual_columns[0]:
                    item_id_value = st.text_input(
                        "Mark / door number",
                        value=clean_issue_mark_for_workbook(issue.get("mark", "")),
                        key=f"{key_prefix}-manual-id",
                    )
                with manual_columns[1]:
                    field_name = st.selectbox(
                        "Workbook field",
                        tuple(WORKBOOK_FIELD_ROWS[tab_name]["fields"]),
                        key=f"{key_prefix}-manual-field",
                    )
                with manual_columns[2]:
                    overwrite_existing = st.checkbox(
                        "Overwrite existing value",
                        key=f"{key_prefix}-manual-overwrite",
                    )
                manual_value = st.text_input(
                    "Verified value to port into Excel",
                    key=f"{key_prefix}-manual-value",
                )
                if st.button(
                    "Port verified value into Excel",
                    disabled=not (
                        item_id_value.strip()
                        and field_name
                        and manual_value.strip()
                    ),
                    key=f"{key_prefix}-manual-apply",
                ):
                    try:
                        workbook_path, workbook_cell = apply_manual_workbook_entry(
                            job_dir,
                            tab_name,
                            item_id_value,
                            field_name,
                            manual_value,
                            overwrite_existing,
                        )
                    except (FileNotFoundError, ValueError, OSError) as exc:
                        st.error(f"Could not port value into Excel: {exc}")
                    else:
                        review_state["issues"][current_issue_id] = {
                            "status": "Resolved",
                            "estimator_note": _join_note_parts(
                                estimator_note,
                                f"Manual workbook entry: {workbook_cell} = {manual_value.strip()}",
                            ),
                            "resolved_by": resolved_by.strip(),
                            "resolution_date": resolution_date.strip()
                            or datetime.now().date().isoformat(),
                        }
                        review_state.setdefault("manual_entries", []).append(
                            {
                                "issue_id": current_issue_id,
                                "applied_at": datetime.now().isoformat(timespec="seconds"),
                                "workbook_path": str(workbook_path.resolve()),
                                "workbook_cell": workbook_cell,
                                "value": manual_value.strip(),
                                "estimator": resolved_by.strip(),
                            }
                        )
                        review_state["approved_for_send"] = False
                        review_state["approved_by"] = ""
                        review_state["approved_at"] = ""
                        save_review_state(job_dir, review_state)
                        _write_job_zip(job_dir)
                        st.success(f"Value ported into `{workbook_cell}` and issue marked Resolved.")
                        st.rerun()
    else:
        st.success("No CRITICAL issues recorded.")

    st.subheader("Estimator QA Resolution Workflow")
    st.caption(
        "Update issue status, add estimator notes, and save. Accepted Risk is recorded explicitly "
        "and removes an issue from the open-CRITICAL gate."
    )
    editable_rows = []
    for issue in qa_issues:
        issue_state = review_state["issues"][issue_id(issue)]
        editable_rows.append(
            {
                "Issue ID": issue_id(issue),
                "Severity": issue.get("severity", ""),
                "Issue type": issue.get("issue_type", ""),
                "Mark": issue.get("mark", ""),
                "Sheet/source page": issue.get("source_page", ""),
                "Description": issue.get("description", ""),
                "Recommended action": issue.get("recommended_action", ""),
                "Status": issue_state["status"],
                "Estimator note": issue_state["estimator_note"],
                "Resolved by": issue_state["resolved_by"],
                "Resolution date": issue_state["resolution_date"],
            }
        )
    edited_rows = st.data_editor(
        editable_rows,
        hide_index=True,
        width="stretch",
        disabled=[
            "Issue ID",
            "Severity",
            "Issue type",
            "Mark",
            "Sheet/source page",
            "Description",
            "Recommended action",
        ],
        column_config={
            "Status": st.column_config.SelectboxColumn(
                "Status", options=REVIEW_STATUSES, required=True
            ),
        },
        key=f"{job_dir.name}-qa-editor",
    )
    if st.button("Save QA resolutions", key=f"{job_dir.name}-save-resolutions"):
        for row in edited_rows:
            review_state["issues"][row["Issue ID"]] = {
                "status": row["Status"],
                "estimator_note": row["Estimator note"],
                "resolved_by": row["Resolved by"],
                "resolution_date": row["Resolution date"],
            }
        review_state["approved_for_send"] = False
        review_state["approved_by"] = ""
        review_state["approved_at"] = ""
        save_review_state(job_dir, review_state)
        _write_job_zip(job_dir)
        st.success("QA resolutions saved locally. Estimator approval was reset for review.")
        st.rerun()

    st.subheader("Source Sheet Previews")
    st.caption("First-page previews from the copied source PDFs. Use them to orient review before opening the originals.")
    preview_paths = sorted((job_dir / "01_Extracted_Pages").glob("*.pdf"))
    preview_columns = st.columns(2)
    for index, pdf_path in enumerate(preview_paths):
        with preview_columns[index % len(preview_columns)]:
            st.image(ensure_pdf_preview(pdf_path), caption=pdf_path.name, width="stretch")
            st.download_button(
                f"Open source PDF: {pdf_path.name}",
                data=pdf_path.read_bytes(),
                file_name=pdf_path.name,
                mime="application/pdf",
                key=f"{job_dir.name}-source-{pdf_path.name}",
                width="stretch",
            )

    st.subheader("Estimator Approval Gate")
    if open_critical_count:
        st.error(f"Approval locked: resolve or explicitly accept risk for {open_critical_count} open CRITICAL issues.")
        st.subheader("Estimator Risk Acceptance")
        if risk_accepted:
            st.success(
                "Risk accepted. Package unlocked for estimator review. "
                "Quote is NOT automatically final."
            )
            st.caption(
                f"Recorded at `{approval_decision.get('approved_at', '')}` with "
                f"`{approval_decision.get('critical_issue_count', open_critical_count)}` "
                "CRITICAL issues."
            )
        qa_report_reviewed = st.checkbox(
            "I opened and reviewed the QA report",
            key=f"{job_dir.name}-risk-qa-reviewed",
        )
        psf_report_reviewed = st.checkbox(
            "I opened and reviewed the PSF/zone report",
            key=f"{job_dir.name}-risk-psf-reviewed",
        )
        unresolved_understood = st.checkbox(
            "I understand this quote still has unresolved CRITICAL issues",
            key=f"{job_dir.name}-risk-critical-understood",
        )
        risk_unlock_text = st.text_input(
            "Type ACCEPT RISK to unlock",
            key=f"{job_dir.name}-risk-unlock-text",
        )
        risk_notes = st.text_area(
            "Reason for accepting risk / estimator notes",
            value=str(approval_decision.get("estimator_notes", "")),
            key=f"{job_dir.name}-risk-notes",
        )
        risk_acceptance_ready = (
            qa_report_reviewed
            and psf_report_reviewed
            and unresolved_understood
            and risk_unlock_text == "ACCEPT RISK"
            and bool(risk_notes.strip())
        )
        if st.button(
            "Accept Risk and Unlock Review Package",
            disabled=not risk_acceptance_ready,
            key=f"{job_dir.name}-accept-risk",
        ) and risk_acceptance_ready:
            save_risk_acceptance(job_dir, open_critical_count, risk_notes)
            _write_job_zip(job_dir)
            st.success(
                "Risk accepted. Package unlocked for estimator review. "
                "Quote is NOT automatically final."
            )
            st.rerun()
    elif review_state["approved_for_send"]:
        st.success(
            f"Approved for send by {review_state['approved_by']} at {review_state['approved_at']}."
        )
    else:
        approval_name = st.text_input(
            "Estimator approval name", key=f"{job_dir.name}-approval-name"
        )
        approval_confirmed = st.checkbox(
            "I reviewed the resolved CRITICAL issues and approve this quote for send.",
            key=f"{job_dir.name}-approval-confirmed",
        )
        if st.button(
            "Record estimator approval",
            disabled=not (approval_name.strip() and approval_confirmed),
            key=f"{job_dir.name}-record-approval",
        ):
            review_state["approved_for_send"] = True
            review_state["approved_by"] = approval_name.strip()
            review_state["approved_at"] = datetime.now().isoformat(timespec="seconds")
            save_review_state(job_dir, review_state)
            _write_job_zip(job_dir)
            st.success("Estimator approval recorded locally.")
            st.rerun()

    st.subheader("Review Order Panel")
    review_steps = (
        "Open QA report",
        "Open PSF/zone report",
        "Open extracted schedules",
        "Open organized plan compartments",
        "Open filled quote workbook",
        "Review proposal draft",
    )
    for number, step in enumerate(review_steps, start=1):
        st.checkbox(f"{number}. {step}", key=f"{job_dir.name}-review-{number}")

    st.subheader("Download Buttons")
    review_files = selected_summary["review_files"]
    download_groups = (
        (
            "Quote And QA Files",
            (
                "filled_quote_workbook.xlsx",
                "qa_report.xlsx",
                "psf_zone_report.xlsx",
            ),
        ),
        (
            "Extracted Schedule CSVs",
            (
                "windows_schedule.csv",
                "storefront_schedule.csv",
                "doors_schedule_candidates.csv",
                "glazing_schedule_candidates.csv",
            ),
        ),
        (
            "Structured Extraction Files",
            (
                "extracted_schedules.json",
                "extraction_audit.json",
                "uncertain_schedule_rows.csv",
                "window_workbook_transfer_audit.json",
                "storefront_workbook_transfer_audit.json",
                "glazing_thermal_requirements.json",
                "door_workbook_transfer_audit.json",
            ),
        ),
        (
            "Organized PDF Packages",
            (
                "floor_plans.zip",
                "architectural_elevations.zip",
                "wind_pressure_elevations.zip",
                "door_schedules.zip",
                "window_schedules.zip",
                "storefront_schedules.zip",
                "organized_plan_set_pdfs.zip",
                "other_plan_documents.zip",
                "source_file_manifest.csv",
                "sheet_index.csv",
            ),
        ),
        (
            "Final Package",
            (
                "proposal_email_draft.txt",
                "full_job_package.zip",
            ),
        ),
    )
    rendered_labels: set[str] = set()
    for group_name, labels in download_groups:
        available_labels = [label for label in labels if label in review_files]
        if not available_labels:
            continue
        st.write(f"**{group_name}**")
        download_columns = st.columns(3)
        for index, label in enumerate(available_labels):
            rendered_labels.add(label)
            display_label = (
                "Review Package - Not Final"
                if label == "full_job_package.zip" and open_critical_count
                else label
            )
            with download_columns[index % len(download_columns)]:
                render_download(
                    job_dir,
                    display_label,
                    review_files[label],
                )
    remaining_downloads = [
        (label, relative_path)
        for label, relative_path in review_files.items()
        if label not in rendered_labels
    ]
    if remaining_downloads:
        st.write("**Other Review Files**")
        download_columns = st.columns(3)
        for index, (label, relative_path) in enumerate(remaining_downloads):
            with download_columns[index % len(download_columns)]:
                render_download(job_dir, label, relative_path)

st.divider()
st.subheader("Job History")
history = load_job_history()
if history:
    st.dataframe(history, hide_index=True, width="stretch")
else:
    st.info("No indexed jobs yet. New dashboard runs are added to jobs_index.csv.")
